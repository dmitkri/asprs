import os
import sqlite3
import secrets
from datetime import datetime, timedelta, timezone, time as dt_time

MOSCOW_TZ = timezone(timedelta(hours=3))
import qrcode
import io
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, BotCommand
from telegram.ext import Application, CommandHandler, CallbackQueryHandler, ContextTypes, MessageHandler, filters
from telegram.error import BadRequest, NetworkError, TimedOut, Conflict

try:
    from config import TELEGRAM_BOT_TOKEN, DATABASE_PATH, MOSCOW_TZ as CONFIG_MOSCOW_TZ
    if CONFIG_MOSCOW_TZ:
        MOSCOW_TZ = CONFIG_MOSCOW_TZ
except ImportError:
    TELEGRAM_BOT_TOKEN = os.getenv('TELEGRAM_BOT_TOKEN')
    DATABASE_PATH = 'aspirs.db'

TELEGRAM_AVAILABLE = True

def get_db():
    conn = sqlite3.connect(DATABASE_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def find_student_by_room_and_info(building, entrance, room_number, fio, phone):
    with get_db() as conn:
        cursor = conn.cursor()
        phone_normalized = ''.join(filter(str.isdigit, phone)) if phone else ''
        try:
            cursor.execute("""
                SELECT id, fio, phone, group_name, building, entrance, room_number, tg_username
                FROM employees 
                WHERE building = ? 
                AND entrance = ? 
                AND room_number = ?
                AND LOWER(fio) LIKE LOWER(?)
                AND REPLACE(REPLACE(REPLACE(REPLACE(phone, ' ', ''), '-', ''), '(', ''), ')', '') LIKE ?
                LIMIT 1
            """, (building, entrance, room_number, f'%{fio}%', f'%{phone_normalized}%'))
        except sqlite3.OperationalError:
            cursor.execute("""
                SELECT id, fio, phone, group_name, building, entrance, room_number
                FROM employees 
                WHERE building = ? 
                AND entrance = ? 
                AND room_number = ?
                AND LOWER(fio) LIKE LOWER(?)
                AND REPLACE(REPLACE(REPLACE(REPLACE(phone, ' ', ''), '-', ''), '(', ''), ')', '') LIKE ?
                LIMIT 1
            """, (building, entrance, room_number, f'%{fio}%', f'%{phone_normalized}%'))
        return cursor.fetchone()

def find_local_student_by_info(fio, phone):
    """Поиск местного студента только по ФИО и телефону (без корпуса/подъезда/комнаты)"""
    with get_db() as conn:
        cursor = conn.cursor()
        # Нормализуем телефон - оставляем только цифры
        phone_normalized = ''.join(filter(str.isdigit, phone)) if phone else ''
        fio_clean = fio.strip() if fio else ''
        
        if not fio_clean:
            return None
        
        # Проверяем наличие поля is_local
        cursor.execute("PRAGMA table_info(employees)")
        columns = [col[1] for col in cursor.fetchall()]
        has_is_local = 'is_local' in columns
        
        # Формируем условие для местных студентов
        # Местный студент - это либо is_local = 1, либо все три поля пустые
        if has_is_local:
            local_condition = "(is_local = 1 OR ((COALESCE(building, '') = '' OR building = '0') AND (COALESCE(entrance, '') = '' OR entrance = '0') AND (COALESCE(room_number, '') = '' OR room_number = '0')))"
        else:
            local_condition = "((COALESCE(building, '') = '' OR building = '0') AND (COALESCE(entrance, '') = '' OR entrance = '0') AND (COALESCE(room_number, '') = '' OR room_number = '0'))"
        
        try:
            # Ищем по ФИО (частичное совпадение) и проверяем телефон в Python
            query = f"""
                SELECT id, fio, phone, group_name, building, entrance, room_number, tg_username
                FROM employees 
                WHERE {local_condition}
                AND LOWER(TRIM(fio)) LIKE LOWER(?)
            """
            cursor.execute(query, (f'%{fio_clean}%',))
            results = cursor.fetchall()
            
            if not results:
                return None
            
            # Функция для нормализации телефона
            def normalize_phone(phone_str):
                if not phone_str:
                    return ''
                return ''.join(filter(str.isdigit, str(phone_str)))
            
            # Если телефон не указан, возвращаем первое совпадение по ФИО
            if not phone_normalized:
                return results[0]
            
            # Ищем совпадение по телефону
            for row in results:
                db_phone = normalize_phone(row['phone'])
                if db_phone:
                    # Проверяем частичное совпадение (на случай разных форматов)
                    if phone_normalized in db_phone or db_phone in phone_normalized:
                        return row
                    # Проверяем последние 10 цифр (обычно это основная часть номера)
                    if len(phone_normalized) >= 10 and len(db_phone) >= 10:
                        if phone_normalized[-10:] == db_phone[-10:]:
                            return row
                else:
                    # Если в базе нет телефона, но ФИО совпадает, возвращаем
                    return row
            
            # Если не нашли по телефону, но есть совпадение по ФИО, возвращаем первое
            return results[0] if results else None
            
        except sqlite3.OperationalError as e:
            # Fallback для старых версий схемы БД
            query = f"""
                SELECT id, fio, phone, group_name, building, entrance, room_number
                FROM employees 
                WHERE {local_condition}
                AND LOWER(TRIM(fio)) LIKE LOWER(?)
            """
            cursor.execute(query, (f'%{fio_clean}%',))
            return cursor.fetchone()

def check_username_match(employee_id, tg_username):
    with get_db() as conn:
        cursor = conn.cursor()
        try:
            cursor.execute("SELECT tg_username FROM employees WHERE id = ?", (employee_id,))
        except sqlite3.OperationalError:
            return True
        student = cursor.fetchone()
        if not student:
            return False
        try:
            student_username = student['tg_username']
        except (KeyError, IndexError):
            return True
        if not student_username or (isinstance(student_username, str) and student_username.strip() == ''):
            return True
        if not tg_username:
            return False
        normalized_username = tg_username.lstrip('@').lower()
        student_username_normalized = student_username.lstrip('@').lower()
        return normalized_username == student_username_normalized

def is_local_student(employee_id):
    """Проверяет, является ли студент местным"""
    with get_db() as conn:
        cursor = conn.cursor()
        # Проверяем наличие поля is_local
        cursor.execute("PRAGMA table_info(employees)")
        columns = [col[1] for col in cursor.fetchall()]
        has_is_local = 'is_local' in columns
        
        if has_is_local:
            cursor.execute("SELECT is_local, building, entrance, room_number FROM employees WHERE id = ?", (employee_id,))
        else:
            cursor.execute("SELECT building, entrance, room_number FROM employees WHERE id = ?", (employee_id,))
        student = cursor.fetchone()
        
        if not student:
            return False
        
        # Если есть поле is_local, используем его как основной источник истины
        if has_is_local:
            try:
                is_local_value = student['is_local']
                # Если is_local явно установлен в 1, то студент местный
                if is_local_value is not None and is_local_value == 1:
                    return True
                # Если is_local = 0 или NULL, то студент НЕ местный (даже если поля пустые)
                return False
            except (KeyError, IndexError):
                # Если не удалось прочитать is_local, считаем что студент НЕ местный
                return False
        
        # Если поля is_local нет в таблице (старая версия БД), 
        # тогда проверяем по пустым полям (обратная совместимость)
        try:
            building = student['building']
        except (KeyError, IndexError):
            building = None
            
        try:
            entrance = student['entrance']
        except (KeyError, IndexError):
            entrance = None
            
        try:
            room_number = student['room_number']
        except (KeyError, IndexError):
            room_number = None
        
        building_empty = not building or building == '' or building == '0'
        entrance_empty = not entrance or entrance == '' or entrance == '0'
        room_empty = not room_number or room_number == '' or room_number == '0'
        
        # Только если поля is_local нет в таблице, используем проверку по пустым полям
        return building_empty and entrance_empty and room_empty

def register_tg_user(tg_user_id, tg_username, employee_id):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT OR REPLACE INTO tg_users (tg_user_id, tg_username, employee_id)
            VALUES (?, ?, ?)
        """, (tg_user_id, tg_username, employee_id))
        conn.commit()

def get_tg_user(tg_user_id):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT tg_user_id, tg_username, employee_id
            FROM tg_users
            WHERE tg_user_id = ?
        """, (tg_user_id,))
        return cursor.fetchone()

def get_all_tg_users():
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT tg_user_id, tg_username, employee_id
            FROM tg_users
            WHERE employee_id IS NOT NULL
        """)
        return cursor.fetchall()

def generate_qr_code(employee_id, tg_user_id):
    with get_db() as conn:
        cursor = conn.cursor()
        qr_code = secrets.token_urlsafe(16)
        now = datetime.now(MOSCOW_TZ)
        expires_at = now + timedelta(minutes=5)
        cursor.execute("""
            INSERT INTO bed_linen_qr_codes (qr_code, employee_id, tg_user_id, created_at, expires_at)
            VALUES (?, ?, ?, ?, ?)
        """, (qr_code, employee_id, tg_user_id, now.isoformat(), expires_at.isoformat()))
        
        conn.commit()
        return qr_code

def cleanup_expired_qr_codes():
    with get_db() as conn:
        cursor = conn.cursor()
        now = datetime.now(MOSCOW_TZ).isoformat()
        cursor.execute("DELETE FROM bed_linen_qr_codes WHERE expires_at < ?", (now,))
        conn.commit()

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tg_user_id = update.effective_user.id
    tg_username = update.effective_user.username or ''
    
    if update.message and update.message.text:
        command_parts = update.message.text.split(' ', 1)
        if len(command_parts) > 1 and command_parts[1].startswith('event_'):
            token = command_parts[1].replace('event_', '')
            await handle_event_token(update, context, tg_user_id, token)
            return
    
    tg_user = get_tg_user(tg_user_id)
    if tg_user and tg_user['employee_id']:
        await show_main_menu(update, context)
    else:
        context.user_data['registration_step'] = 'building'
        keyboard = [[InlineKeyboardButton("Я местный!", callback_data="local_student")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        await update.message.reply_text(
            "Привет! Я твой помощник в кампусе.\n\n"
            "Давай зарегистрируемся! Мне нужна небольшая информация:\n\n"
            "Отправь номер корпуса:",
            reply_markup=reply_markup
        )

async def handle_event_token(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id, token):
    tg_username = update.effective_user.username or ''
    if not tg_username:
        await update.message.reply_text(
            "Для регистрации студентов на мероприятия нужно указать username в Telegram."
        )
        return
    
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT employee_id, expires_at FROM student_event_tokens 
            WHERE token = ?
        """, (token,))
        token_data = cursor.fetchone()
        
        if not token_data:
            await update.message.reply_text(
            "Упс! Неверный токен QR-кода.\n"
            "Убедись, что сканируешь актуальный QR-код студента."
            )
            return
        expires_at = datetime.fromisoformat(token_data['expires_at'])
        if expires_at.tzinfo is None:
            expires_at = expires_at.replace(tzinfo=MOSCOW_TZ)
        else:
            expires_at = expires_at.astimezone(MOSCOW_TZ)
        
        if datetime.now(MOSCOW_TZ) > expires_at:
            await update.message.reply_text(
            "Ой! Токен QR-кода истек.\n"
            "Студенту нужно получить новый QR-код."
            )
            return
        
        employee_id = token_data['employee_id']
        is_global_org = cursor.execute(
            'SELECT id FROM event_organizers WHERE tg_username = ? OR tg_user_id = ?',
            (tg_username.lstrip('@'), tg_user_id)
        ).fetchone()
        
        if not is_global_org:
            await update.message.reply_text(
            "К сожалению, у тебя нет прав организатора мероприятий.\n"
            "Обратись к администратору для получения прав."
            )
            return
        cursor.execute("SELECT id, fio FROM employees WHERE id = ?", (employee_id,))
        student = cursor.fetchone()
        
        if not student:
            await update.message.reply_text("Упс! Студент не найден в базе данных.")
            return
        today = datetime.now(MOSCOW_TZ).date().isoformat()
        events = cursor.execute("""
            SELECT id, location, description, event_date, event_time, organizers
            FROM events
            WHERE event_date >= ?
            ORDER BY event_date, event_time
            LIMIT 20
        """, (today,)).fetchall()
        
        if not events:
            await update.message.reply_text(
                f"Студент: {student['fio']}\n\n"
                "Пока нет активных мероприятий для регистрации."
            )
            return
        context.user_data['event_token'] = token
        context.user_data['event_student_id'] = employee_id
        context.user_data['event_student_fio'] = student['fio']
        context.user_data['event_step'] = 'select_event'
        
        keyboard = []
        for event in events:
            event_date = event['event_date']
            event_time = event['event_time']
            location = event['location'][:30] if event['location'] else 'Без места'
            button_text = f"{event_date} {event_time} - {location}"
            keyboard.append([InlineKeyboardButton(button_text, callback_data=f"register_event_{event['id']}")])
        keyboard.append([InlineKeyboardButton("Отмена", callback_data="back_to_menu")])
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        message = (
            f"Студент: {student['fio']}\n\n"
            "Выбери мероприятие для регистрации:"
        )
        
        await update.message.reply_text(message, reply_markup=reply_markup)

async def qr_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tg_user_id = update.effective_user.id
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        await update.message.reply_text(
            "Ты еще не зарегистрирован! Используй /start для регистрации."
        )
        return
    await generate_qr_for_user(update, context, tg_user_id)

async def menu_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tg_user_id = update.effective_user.id
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'message') and update.message:
            await update.message.reply_text(
                "Ты еще не зарегистрирован!\n\n"
                "Используй команду /start для регистрации."
            )
        elif hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Ты еще не зарегистрирован! Используй /start для регистрации.", show_alert=True)
        return
    await show_main_menu(update, context)

async def show_main_menu(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tg_user_id = update.effective_user.id
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'message') and update.message:
            await update.message.reply_text("Ты еще не зарегистрирован! Используй /start для регистрации.")
        elif hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.message.reply_text("Ты еще не зарегистрирован! Используй /start для регистрации.")
            return
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT fio, group_name, building, entrance, room_number FROM employees WHERE id = ?", (tg_user['employee_id'],))
        student = cursor.fetchone()

        tg_username = update.effective_user.username or ''
        is_organizer = False
        is_elder = False
        is_chief = False
        
        # Проверяем, является ли студент местным (используем единую функцию)
        is_local = is_local_student(tg_user['employee_id'])
        
        # Проверка организатора (требует username)
        if tg_username:
            cursor.execute("SELECT id FROM event_organizers WHERE tg_username = ? OR tg_user_id = ?", 
                         (tg_username.lstrip('@'), tg_user_id))
            is_organizer = cursor.fetchone() is not None
        
        # Проверка старосты подъезда (не требует username)
        cursor.execute("""
            SELECT ea.id FROM elder_assignments ea
            JOIN employees e ON ea.employee_id = e.id
            WHERE e.id = ? AND e.is_entrance_elder = 1
        """, (tg_user['employee_id'],))
        is_elder = cursor.fetchone() is not None
        
        # Проверка старшего студента (ответственного за обходы)
        # Проверяем и по полю is_round_chief, и по наличию записи в round_assignments
        cursor.execute("""
            SELECT is_round_chief FROM employees WHERE id = ?
        """, (tg_user['employee_id'],))
        student_info = cursor.fetchone()
        try:
            if student_info and student_info['is_round_chief'] == 1:
                is_chief = True
            else:
                # Если is_round_chief не установлен, проверяем наличие записи в round_assignments
                cursor.execute("""
                    SELECT id FROM round_assignments 
                    WHERE student_id = ? AND building IS NOT NULL AND entrance IS NOT NULL
                """, (tg_user['employee_id'],))
                is_chief = cursor.fetchone() is not None
        except (KeyError, IndexError, TypeError):
            # Если поле is_round_chief отсутствует или не может быть прочитано, проверяем round_assignments
            cursor.execute("""
                SELECT id FROM round_assignments 
                WHERE student_id = ? AND building IS NOT NULL AND entrance IS NOT NULL
            """, (tg_user['employee_id'],))
            is_chief = cursor.fetchone() is not None
    
    if student:
        # Для местных студентов показываем только QR для мероприятий
        if is_local:
            keyboard = [
                [InlineKeyboardButton("Мой QR-код для мероприятий", callback_data="get_event_qr")]
            ]
            
            if is_organizer:
                keyboard.append([InlineKeyboardButton("Сканировать QR на мероприятии", callback_data="scan_event")])
            
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            menu_text = f"Привет, {student['fio']}!\n\n"
            menu_text += "Что хочешь сделать?\n"
            menu_text += "• Мой QR-код для мероприятий\n"
            if is_organizer:
                menu_text += "• Сканировать QR на мероприятии\n"
        else:
            keyboard = [
                [InlineKeyboardButton("Получить QR-код для белья", callback_data="get_qr")],
                [InlineKeyboardButton("Мой QR-код для мероприятий", callback_data="get_event_qr")],
                [InlineKeyboardButton("Оплатить проживание", callback_data="pay_dormitory")],
                [InlineKeyboardButton("Подать заявку", callback_data="create_maintenance")],
                [InlineKeyboardButton("Мои заявки", callback_data="my_maintenance")]
            ]
            
            if is_organizer:
                keyboard.append([InlineKeyboardButton("Сканировать QR на мероприятии", callback_data="scan_event")])
            if is_elder:
                keyboard.append([InlineKeyboardButton("График дежурств", callback_data="duty_schedule")])
            if is_chief:
                keyboard.append([InlineKeyboardButton("Назначить на обход", callback_data="assign_round")])
            
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            menu_text = f"Привет, {student['fio']}!\n\n"
            menu_text += "Что хочешь сделать?\n"
            menu_text += "• Получить QR-код для белья\n"
            menu_text += "• Мой QR-код для мероприятий\n"
            menu_text += "• Оплатить проживание\n"
            menu_text += "• Подать заявку\n"
            menu_text += "• Мои заявки\n"
            if is_organizer:
                menu_text += "• Сканировать QR на мероприятии\n"
            if is_elder:
                menu_text += "• График дежурств\n"
            if is_chief:
                menu_text += "• Назначить на обход\n"
        
        if hasattr(update, 'message') and update.message:
            await update.message.reply_text(
                menu_text,
                reply_markup=reply_markup
            )
        elif hasattr(update, 'callback_query') and update.callback_query:
            try:
                message = update.callback_query.message
                if message.text:
                    await update.callback_query.message.edit_text(
                        menu_text,
                        reply_markup=reply_markup
                    )
                else:
                    await update.callback_query.message.reply_text(
                        menu_text,
                        reply_markup=reply_markup
                    )
            except BadRequest as e:
                if "Message is not modified" in str(e):
                    pass
                else:
                    await update.callback_query.message.reply_text(
                        menu_text,
                        reply_markup=reply_markup
                    )
            except Exception as e:
                await update.callback_query.message.reply_text(
                    menu_text,
                    reply_markup=reply_markup
                )

async def button_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        query = update.callback_query
        if not query:
            return
        
        tg_user_id = query.from_user.id
        tg_user = get_tg_user(tg_user_id)
        
        # Проверяем, является ли студент местным (для блокировки недоступных функций)
        is_local = False
        if tg_user and tg_user['employee_id']:
            is_local = is_local_student(tg_user['employee_id'])
        
        if query.data == "get_qr":
            if is_local:
                await query.answer("Эта функция недоступна для местных студентов.", show_alert=True)
                return
            await generate_qr_for_user(update, context, tg_user_id)
        elif query.data == "pay_dormitory":
            if is_local:
                await query.answer("Эта функция недоступна для местных студентов.", show_alert=True)
                return
            await generate_payment_qr(update, context, tg_user_id)
        elif query.data == "get_event_qr":
            await generate_event_qr(update, context, tg_user_id)
        elif query.data == "create_maintenance":
            if is_local:
                await query.answer("Эта функция недоступна для местных студентов.", show_alert=True)
                return
            await create_maintenance_request(update, context, tg_user_id)
        elif query.data == "my_maintenance":
            if is_local:
                await query.answer("Эта функция недоступна для местных студентов.", show_alert=True)
                return
            await show_my_maintenance_requests(update, context, tg_user_id)
        elif query.data == "scan_event":
            await scan_event_qr(update, context, tg_user_id)
        elif query.data.startswith("select_event_"):
            event_id = int(query.data.replace("select_event_", ""))
            context.user_data['selected_event_id'] = event_id
            context.user_data['scan_event_step'] = 'waiting_qr'
            await query.answer()
            message = (
                "Сканирование QR-кода\n\n"
                "Используй камеру телефона для сканирования QR-кода."
            )
            keyboard = [[InlineKeyboardButton("Отмена", callback_data="back_to_menu")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            await query.message.reply_text(message, reply_markup=reply_markup)
        elif query.data.startswith("register_event_"):
            event_id = int(query.data.replace("register_event_", ""))
            await register_student_to_event(update, context, tg_user_id, event_id)
        elif query.data == "duty_schedule":
            await show_duty_schedule(update, context, tg_user_id)
        elif query.data == "generate_duty_current":
            now = datetime.now(MOSCOW_TZ)
            await generate_duty_schedule(update, context, tg_user_id, now.month, now.year)
        elif query.data == "generate_duty_next":
            now = datetime.now(MOSCOW_TZ)
            if now.month == 12:
                next_month = 1
                next_year = now.year + 1
            else:
                next_month = now.month + 1
                next_year = now.year
            await generate_duty_schedule(update, context, tg_user_id, next_month, next_year)
        elif query.data == "assign_round":
            await assign_round(update, context, tg_user_id)
        elif query.data.startswith("assign_round_student_"):
            student_id = int(query.data.replace("assign_round_student_", ""))
            await handle_assign_round_student(update, context, tg_user_id, student_id)
        elif query.data.startswith("assign_round_building_"):
            parts = query.data.replace("assign_round_building_", "").split("|")
            if len(parts) == 3:
                building = parts[0]
                entrance = parts[1]
                student_id = int(parts[2])
                await handle_assign_round_final(update, context, tg_user_id, student_id, building, entrance)
        elif query.data == "back_to_menu":
            # Очищаем состояние регистрации/заявки при возврате в меню
            context.user_data.pop('registration_step', None)
            context.user_data.pop('maintenance_step', None)
            await show_main_menu(update, context)
        elif query.data == "local_student":
            await query.answer()
            context.user_data['registration_step'] = 'local_fio'
            context.user_data['is_local'] = True
            await query.message.reply_text(
                "Отлично! Ты местный студент.\n\n"
                "Пожалуйста, отправь свое ФИО:"
            )
        else:
            await query.answer("Упс! Неизвестная команда.", show_alert=True)
    except Exception as e:
        import traceback
        traceback.print_exc()
        pass

async def error_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    import traceback
    import logging
    
    logger = logging.getLogger(__name__)
    
    if context.error and isinstance(context.error, Conflict):
        error_msg = str(context.error)
        if 'getUpdates' in error_msg or 'only one bot instance' in error_msg.lower():
            logger.error("Обнаружено несколько экземпляров бота! Остановите другие процессы.")
            logger.error("Для проверки запущенных процессов используйте: ps aux | grep telegram_bot")
            logger.error("Для остановки процесса используйте: kill <PID>")
            return
    
    logger.error(f"Exception while handling an update: {context.error}", exc_info=context.error)
    
    if NetworkError and isinstance(context.error, NetworkError):
        error_message = str(context.error)
        return
    if TimedOut and isinstance(context.error, TimedOut):
        error_message = str(context.error)
        return
    if BadRequest and isinstance(context.error, BadRequest):
        error_message = str(context.error)
        if "no text in the message" in error_message.lower():
            return
    
    if update and update.effective_chat:
        try:
            await context.bot.send_message(
                chat_id=update.effective_chat.id,
                text="Упс! Что-то пошло не так. Попробуй позже или используй команду /start"
            )
        except Exception as e:
            pass

def check_active_period():
    conn = get_db()
    cursor = conn.cursor()
    today_str = datetime.now(MOSCOW_TZ).date().isoformat()
    cursor.execute("""
        SELECT name FROM sqlite_master 
        WHERE type='table' AND name='bed_linen_dates'
    """)
    table_exists = cursor.fetchone()
    if not table_exists:
        conn.close()
        return False
    cursor.execute("""
        SELECT id FROM bed_linen_dates
        WHERE is_active = 1 AND start_date <= ? AND end_date >= ?
        ORDER BY start_date DESC
        LIMIT 1
    """, (today_str, today_str))
    period = cursor.fetchone()
    conn.close()
    result = period is not None
    return result

def get_active_period():
    conn = get_db()
    cursor = conn.cursor()
    today_str = datetime.now(MOSCOW_TZ).date().isoformat()
    cursor.execute("""
        SELECT name FROM sqlite_master 
        WHERE type='table' AND name='bed_linen_dates'
    """)
    table_exists = cursor.fetchone()
    if not table_exists:
        conn.close()
        return None
    cursor.execute("""
        SELECT id, start_date, end_date
        FROM bed_linen_dates
        WHERE is_active = 1 AND start_date <= ? AND end_date >= ?
        ORDER BY start_date DESC
        LIMIT 1
    """, (today_str, today_str))
    period = cursor.fetchone()
    conn.close()
    if period:
        return {
            'id': period['id'],
            'start_date': period['start_date'],
            'end_date': period['end_date']
        }
    return None

def get_next_period():
    conn = get_db()
    cursor = conn.cursor()
    today_str = datetime.now(MOSCOW_TZ).date().isoformat()
    cursor.execute("""
        SELECT name FROM sqlite_master 
        WHERE type='table' AND name='bed_linen_dates'
    """)
    table_exists = cursor.fetchone()
    if not table_exists:
        conn.close()
        return None
    cursor.execute("""
        SELECT id, start_date, end_date
        FROM bed_linen_dates
        WHERE is_active = 1 AND start_date > ?
        ORDER BY start_date ASC
        LIMIT 1
    """, (today_str,))
    period = cursor.fetchone()
    conn.close()
    if period:
        return {
            'id': period['id'],
            'start_date': period['start_date'],
            'end_date': period['end_date']
        }
    return None

def has_received_bed_linen(employee_id, period_id):
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT name FROM sqlite_master 
        WHERE type='table' AND name='scans'
    """)
    table_exists = cursor.fetchone()
    if not table_exists:
        conn.close()
        return False
    cursor.execute("""
        SELECT id FROM scans 
        WHERE employee_id = ? AND period_id = ?
        LIMIT 1
    """, (employee_id, period_id))
    scan = cursor.fetchone()
    conn.close()
    return scan is not None

def get_period_starting_tomorrow():
    conn = get_db()
    cursor = conn.cursor()
    tomorrow = (datetime.now(MOSCOW_TZ).date() + timedelta(days=1)).isoformat()
    cursor.execute("""
        SELECT name FROM sqlite_master 
        WHERE type='table' AND name='bed_linen_dates'
    """)
    table_exists = cursor.fetchone()
    if not table_exists:
        conn.close()
        return None
    cursor.execute("""
        SELECT id, start_date, end_date
        FROM bed_linen_dates
        WHERE is_active = 1 AND start_date = ?
        ORDER BY start_date ASC
        LIMIT 1
    """, (tomorrow,))
    period = cursor.fetchone()
    conn.close()
    if period:
        return {
            'id': period['id'],
            'start_date': period['start_date'],
            'end_date': period['end_date']
        }
    return None

async def send_bed_linen_reminders(context: ContextTypes.DEFAULT_TYPE):
    period = get_period_starting_tomorrow()
    if not period:
        return
    try:
        start_date = datetime.strptime(period['start_date'], "%Y-%m-%d").date()
        end_date = datetime.strptime(period['end_date'], "%Y-%m-%d").date()
        start_formatted = start_date.strftime("%d.%m.%Y")
        end_formatted = end_date.strftime("%d.%m.%Y")
        if start_date == end_date:
            period_text = start_formatted
        else:
            period_text = f"с {start_formatted} по {end_formatted}"
    except Exception as e:
        period_text = period['start_date']
    message = (
        "Напоминание о выдаче постельного белья\n\n"
        f"Завтра начинается период выдачи белья:\n{period_text}\n\n"
        "Не забудь получить QR-код для получения белья!\n"
        "Используй команду /qr или кнопку в меню бота."
    )
    users = get_all_tg_users()
    if not users:
        return
    sent_count = 0
    failed_count = 0
    for user in users:
        tg_user_id = user['tg_user_id']
        try:
            await context.bot.send_message(
                chat_id=tg_user_id,
                text=message
            )
            sent_count += 1
        except Exception as e:
            failed_count += 1

async def generate_qr_for_user(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            try:
                await update.callback_query.edit_message_text("Ты еще не зарегистрирован! Используй /start для регистрации.")
            except Exception as e:
                await update.callback_query.message.reply_text("Ты еще не зарегистрирован! Используй /start для регистрации.")
        elif hasattr(update, 'message') and update.message:
            await update.message.reply_text("Ты еще не зарегистрирован! Используй /start для регистрации.")
        return
    
    employee_id = tg_user['employee_id']
    active_period = get_active_period()
    if not active_period:
        next_period = get_next_period()
        error_message = "Сейчас нет активного периода выдачи белья.\n\n"
        error_message += "QR-код можно получить только в дни выдачи, установленные администратором.\n\n"
        if next_period:
            try:
                start_date = datetime.strptime(next_period['start_date'], "%Y-%m-%d").date()
                end_date = datetime.strptime(next_period['end_date'], "%Y-%m-%d").date()
                start_formatted = start_date.strftime("%d.%m.%Y")
                end_formatted = end_date.strftime("%d.%m.%Y")
                if start_date == end_date:
                    error_message += f"Следующий период выдачи белья:\n{start_formatted}"
                else:
                    error_message += f"Следующий период выдачи белья:\nс {start_formatted} по {end_formatted}"
            except Exception as e:
                error_message += "Обратись к администратору для создания периода выдачи белья."
        else:
            error_message += "Обратись к администратору для создания периода выдачи белья."
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer(error_message, show_alert=True)
            try:
                await update.callback_query.message.reply_text(error_message)
            except Exception as e:
                pass
        elif hasattr(update, 'message') and update.message:
            await update.message.reply_text(error_message)
        return
    
    period_id = active_period['id']
    has_received = has_received_bed_linen(employee_id, period_id)
    if has_received:
        next_period = get_next_period()
        error_message = "Ты уже получил комплект постельного белья в этом периоде.\n\n"
        error_message += "Следующая выдача будет в следующем периоде.\n\n"
        if next_period:
            try:
                start_date = datetime.strptime(next_period['start_date'], "%Y-%m-%d").date()
                end_date = datetime.strptime(next_period['end_date'], "%Y-%m-%d").date()
                start_formatted = start_date.strftime("%d.%m.%Y")
                end_formatted = end_date.strftime("%d.%m.%Y")
                if start_date == end_date:
                    error_message += f"Следующий период выдачи белья:\n{start_formatted}"
                else:
                    error_message += f"Следующий период выдачи белья:\nс {start_formatted} по {end_formatted}"
            except Exception as e:
                error_message += "Обратись к администратору для получения информации о следующем периоде."
        else:
            error_message += "Обратись к администратору для получения информации о следующем периоде."
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer(error_message, show_alert=True)
            try:
                await update.callback_query.message.reply_text(error_message)
            except Exception as e:
                pass
        elif hasattr(update, 'message') and update.message:
            await update.message.reply_text(error_message)
        return
    cleanup_expired_qr_codes()
    qr_code = generate_qr_code(employee_id, tg_user_id)
    qr_img = qrcode.make(qr_code)
    qr_buf = io.BytesIO()
    qr_img.save(qr_buf, format='PNG')
    qr_buf.seek(0)
    message_text = (
        "Готово! Вот твой QR-код для получения белья:\n\n"
        f"Код действителен 5 минут\n"
        f"Покажи этот QR-код администратору для сканирования"
    )
    keyboard = [
        [InlineKeyboardButton("Получить новый QR-код", callback_data="get_qr")]
    ]
    reply_markup = InlineKeyboardMarkup(keyboard)
    photo_caption = (
        f"Готово! Вот твой QR-код для получения белья:\n\n"
        f"Код: {qr_code}\n"
        f"Код действителен 5 минут\n"
        f"Действителен до: {datetime.now(MOSCOW_TZ) + timedelta(minutes=5):%H:%M:%S}\n\n"
        f"Покажи этот QR-код администратору для сканирования"
    )
    if hasattr(update, 'callback_query') and update.callback_query:
        try:
            await context.bot.send_photo(
                chat_id=tg_user_id,
                photo=qr_buf,
                caption=photo_caption,
                reply_markup=reply_markup
            )
            try:
                message = update.callback_query.message
                if message.text:
                    await update.callback_query.edit_message_text(
                        message_text,
                        reply_markup=reply_markup
                    )
                elif message.caption:
                    await update.callback_query.edit_message_caption(
                        caption=message_text,
                        reply_markup=reply_markup
                    )
            except Exception as e:
                pass
        except Exception as e:
            pass
    elif hasattr(update, 'message') and update.message:
        await update.message.reply_text(message_text, reply_markup=reply_markup)
        try:
            await context.bot.send_photo(
                chat_id=tg_user_id,
                photo=qr_buf,
                caption=photo_caption,
                reply_markup=reply_markup
            )
        except Exception as e:
            pass

async def generate_payment_qr(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Вы не зарегистрированы. Используйте /start для регистрации.", show_alert=True)
        return
    employee_id = tg_user['employee_id']
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, fio, building, entrance, room_number, room_occupancy FROM employees WHERE id = ?", (employee_id,))
        student = cursor.fetchone()
    if not student:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Студент не найден", show_alert=True)
        return
    try:
        PAYMENT_DETAILS = {
            'bank_name': 'Сбербанк',
            'recipient_name': 'АНОО ВО "Университет "Сириус"',
            'inn': '2367010021',
            'bic': '046015602',
            'account': '40703810530060000441',
            'bank_account': '30101810100000000602',
        }
        
        PAYMENT_RATES = {
            1: 5800,
            2: 2900,
            3: 1500,
            4: 900,
        }
        
        def calculate_payment_amount(room_occupancy):
            if not room_occupancy or room_occupancy <= 0:
                room_occupancy = 1
            
            if room_occupancy > 4:
                room_occupancy = 4
            
            return PAYMENT_RATES.get(room_occupancy, PAYMENT_RATES[1])
        
        student_fio = student['fio'] or 'Неизвестно'
        try:
            student_building = (student['building'] or '').strip() if student['building'] else ''
        except (KeyError, IndexError):
            student_building = ''
        try:
            student_entrance = (student['entrance'] or '').strip() if student['entrance'] else ''
        except (KeyError, IndexError):
            student_entrance = ''
        try:
            student_room = (student['room_number'] or '').strip() if student['room_number'] else ''
        except (KeyError, IndexError):
            student_room = ''
        
        
        room_building = student_building
        room_entrance = student_entrance
        room_number_display = student_room
        
        room_occupancy = None
        
        if student_building and student_entrance and student_room:
            with get_db() as conn:
                try:
                    room = conn.execute(
                        'SELECT building, entrance, room_number, capacity FROM rooms WHERE building = ? AND entrance = ? AND room_number = ?',
                        (student_building, student_entrance, student_room)
                    ).fetchone()
                    
                    
                    if room:
                        try:
                            room_occupancy = room['capacity']
                            room_building = room['building'] if room['building'] else student_building
                            room_entrance = room['entrance'] if room['entrance'] else student_entrance
                            room_number_display = room['room_number'] if room['room_number'] else student_room
                        except (KeyError, IndexError):
                            pass
                    
                    if not room_occupancy:
                        room_full = conn.execute('''
                            SELECT building, entrance, room_number, capacity FROM rooms 
                            WHERE TRIM(building) = ? 
                            AND TRIM(entrance) = ? 
                            AND TRIM(room_number) = ?
                        ''', (student_building.strip(), student_entrance.strip(), student_room.strip())).fetchone()
                        
                        if room_full:
                            try:
                                room_occupancy = room_full['capacity']
                                room_building = room_full['building'] if room_full['building'] else student_building
                                room_entrance = room_full['entrance'] if room_full['entrance'] else student_entrance
                                room_number_display = room_full['room_number'] if room_full['room_number'] else student_room
                            except (KeyError, IndexError):
                                pass
                        
                except sqlite3.OperationalError as e:
                    pass
                except Exception as e:
                    import traceback
                    traceback.print_exc()
                    pass
        
        if not room_occupancy:
            with get_db() as conn:
                try:
                    room = conn.execute('''
                        SELECT DISTINCT r.capacity 
                        FROM rooms r
                        INNER JOIN employees e ON TRIM(e.building) = TRIM(r.building) 
                            AND TRIM(e.entrance) = TRIM(r.entrance) 
                            AND TRIM(e.room_number) = TRIM(r.room_number)
                        WHERE e.id = ?
                    ''', (employee_id,)).fetchone()
                    
                    room_data = conn.execute('''
                        SELECT DISTINCT r.building, r.entrance, r.room_number, r.capacity 
                        FROM rooms r
                        INNER JOIN employees e ON TRIM(COALESCE(e.building, '')) = TRIM(COALESCE(r.building, '')) 
                            AND TRIM(COALESCE(e.entrance, '')) = TRIM(COALESCE(r.entrance, '')) 
                            AND TRIM(COALESCE(e.room_number, '')) = TRIM(COALESCE(r.room_number, ''))
                        WHERE e.id = ?
                    ''', (employee_id,)).fetchone()
                    
                    if room_data:
                        try:
                            if room_data['capacity']:
                                room_occupancy = room_data['capacity']
                                room_building = room_data['building'] if room_data['building'] else student_building
                                room_entrance = room_data['entrance'] if room_data['entrance'] else student_entrance
                                room_number_display = room_data['room_number'] if room_data['room_number'] else student_room
                        except (KeyError, IndexError):
                            pass
                except Exception as e:
                    import traceback
                    traceback.print_exc()
                    pass
        
        if not room_occupancy:
            with get_db() as conn:
                try:
                    
                    rooms_with_student = conn.execute('''
                        SELECT DISTINCT r.building, r.entrance, r.room_number, r.capacity
                        FROM rooms r
                        WHERE EXISTS (
                            SELECT 1 FROM employees e 
                            WHERE e.id = ? 
                            AND (
                                (e.building IS NOT NULL AND e.building != '' AND TRIM(e.building) = TRIM(r.building)) OR e.building IS NULL OR e.building = ''
                            )
                            AND (
                                (e.entrance IS NOT NULL AND e.entrance != '' AND TRIM(e.entrance) = TRIM(r.entrance)) OR e.entrance IS NULL OR e.entrance = ''
                            )
                            AND (
                                (e.room_number IS NOT NULL AND e.room_number != '' AND TRIM(e.room_number) = TRIM(r.room_number)) OR e.room_number IS NULL OR e.room_number = ''
                            )
                        )
                    ''', (employee_id,)).fetchall()
                    
                    if rooms_with_student:
                        room = rooms_with_student[0]
                        room_occupancy = room['capacity']
                        room_building = room['building'] if room['building'] else student_building
                        room_entrance = room['entrance'] if room['entrance'] else student_entrance
                        room_number_display = room['room_number'] if room['room_number'] else student_room
                    
                    if not room_occupancy:
                        all_rooms = conn.execute('SELECT building, entrance, room_number, capacity FROM rooms').fetchall()
                        for room in all_rooms:
                            room_b = (room['building'] or '').strip()
                            room_e = (room['entrance'] or '').strip()
                            room_r = (room['room_number'] or '').strip()
                            
                            students_in_room = conn.execute('''
                                SELECT id, building, entrance, room_number FROM employees 
                                WHERE TRIM(COALESCE(building, '')) = TRIM(COALESCE(?, ''))
                                AND TRIM(COALESCE(entrance, '')) = TRIM(COALESCE(?, ''))
                                AND TRIM(COALESCE(room_number, '')) = TRIM(COALESCE(?, ''))
                            ''', (room_b, room_e, room_r)).fetchall()
                            
                            
                            for emp in students_in_room:
                                if emp['id'] == employee_id:
                                    room_occupancy = room['capacity']
                                    room_building = room_b if room_b else ''
                                    room_entrance = room_e if room_e else ''
                                    room_number_display = room_r if room_r else ''
                                    break
                            
                            if room_occupancy:
                                break
                        
                        if not room_occupancy and (student_building or student_entrance or student_room):
                            for room in all_rooms:
                                room_b = (room['building'] or '').strip().lower()
                                room_e = (room['entrance'] or '').strip().lower()
                                room_r = (room['room_number'] or '').strip().lower()
                                
                                student_b = student_building.strip().lower() if student_building else ''
                                student_e = student_entrance.strip().lower() if student_entrance else ''
                                student_r = student_room.strip().lower() if student_room else ''
                                
                                matches = 0
                                if student_b and room_b == student_b:
                                    matches += 1
                                if student_e and room_e == student_e:
                                    matches += 1
                                if student_r and room_r == student_r:
                                    matches += 1
                                
                                total_fields = sum([1 for x in [student_b, student_e, student_r] if x])
                                if total_fields > 0 and matches >= min(2, total_fields):
                                    room_occupancy = room['capacity']
                                    room_building = room['building'] if room['building'] else student_building
                                    room_entrance = room['entrance'] if room['entrance'] else student_entrance
                                    room_number_display = room['room_number'] if room['room_number'] else student_room
                                    break
                                
                except Exception as e:
                    import traceback
                    traceback.print_exc()
                    pass
        
        if not room_occupancy:
            try:
                room_occupancy = student['room_occupancy'] if 'room_occupancy' in student and student['room_occupancy'] else 1
            except (KeyError, TypeError):
                room_occupancy = 1
        
        if not room_occupancy or room_occupancy == 0:
            room_occupancy = 1
        
        if room_occupancy and room_occupancy > 0 and (not room_building or not room_entrance or not room_number_display):
            with get_db() as conn:
                try:
                    all_rooms_with_capacity = conn.execute('''
                        SELECT r.building, r.entrance, r.room_number, r.capacity
                        FROM rooms r
                        WHERE r.capacity = ?
                    ''', (room_occupancy,)).fetchall()
                    
                    
                    for room_candidate in all_rooms_with_capacity:
                        room_cand_b = (room_candidate['building'] or '').strip()
                        room_cand_e = (room_candidate['entrance'] or '').strip()
                        room_cand_r = (room_candidate['room_number'] or '').strip()
                        
                        student_in_room = conn.execute('''
                            SELECT id FROM employees 
                            WHERE id = ?
                            AND TRIM(COALESCE(building, '')) = TRIM(COALESCE(?, ''))
                            AND TRIM(COALESCE(entrance, '')) = TRIM(COALESCE(?, ''))
                            AND TRIM(COALESCE(room_number, '')) = TRIM(COALESCE(?, ''))
                        ''', (employee_id, room_cand_b, room_cand_e, room_cand_r)).fetchone()
                        
                        if student_in_room:
                            room_building = room_cand_b if room_cand_b else ''
                            room_entrance = room_cand_e if room_cand_e else ''
                            room_number_display = room_cand_r if room_cand_r else ''
                            break
                    
                    if not room_building or not room_entrance or not room_number_display:
                        room_direct = conn.execute('''
                            SELECT r.building, r.entrance, r.room_number
                            FROM rooms r
                            WHERE EXISTS (
                                SELECT 1 FROM employees e
                                WHERE e.id = ?
                                AND (
                                    (TRIM(COALESCE(e.building, '')) = TRIM(COALESCE(r.building, '')) AND TRIM(COALESCE(e.building, '')) != '') OR
                                    (TRIM(COALESCE(e.entrance, '')) = TRIM(COALESCE(r.entrance, '')) AND TRIM(COALESCE(e.entrance, '')) != '') OR
                                    (TRIM(COALESCE(e.room_number, '')) = TRIM(COALESCE(r.room_number, '')) AND TRIM(COALESCE(e.room_number, '')) != '')
                                )
                            )
                            LIMIT 1
                        ''', (employee_id,)).fetchone()
                        
                        if room_direct:
                            room_building = room_direct['building'] if room_direct['building'] else room_building
                            room_entrance = room_direct['entrance'] if room_direct['entrance'] else room_entrance
                            room_number_display = room_direct['room_number'] if room_direct['room_number'] else room_number_display
                    
                except Exception as e:
                    import traceback
                    traceback.print_exc()
        
        if not room_building or room_building == '':
            room_building = student_building if student_building else ''
        if not room_entrance or room_entrance == '':
            room_entrance = student_entrance if student_entrance else ''
        if not room_number_display or room_number_display == '':
            room_number_display = student_room if student_room else ''
        
        
        payment_amount = calculate_payment_amount(room_occupancy)
        
        payment_purpose = f'Оплата за проживание в общежитии Гамма студента {student_fio}.'
        
        qr_string = (
            f"ST00012|Name={PAYMENT_DETAILS['recipient_name']}|"
            f"PersonalAcc={PAYMENT_DETAILS['account']}|"
            f"BankName={PAYMENT_DETAILS['bank_name']}|"
            f"BIC={PAYMENT_DETAILS['bic']}|"
            f"CorrespAcc={PAYMENT_DETAILS['bank_account']}|"
            f"Purpose={payment_purpose}|"
            f"Sum={payment_amount * 100}"
        )
        
        try:
            qr = qrcode.QRCode(version=1, box_size=10, border=5)
            qr.add_data(qr_string)
            qr.make(fit=True)
            qr_img = qr.make_image(fill_color="black", back_color="white")
            
            qr_buf = io.BytesIO()
            qr_img.save(qr_buf, format='PNG')
            qr_buf.seek(0)
            
            room_info = f"{room_building}-{room_entrance}-{room_number_display}" if (room_building and room_entrance and room_number_display) else "не указана"
            payment_text = (
                f"*Оплата проживания*\n\n"
                f"Студент: {student_fio}\n"
                f"Комната: {room_info}\n"
                f"Количество проживающих: {room_occupancy}\n"
                f"*Сумма к оплате: {payment_amount} руб.*\n\n"
                f"*Реквизиты для оплаты:*\n"
                f"Банк: {PAYMENT_DETAILS['bank_name']}\n"
                f"Получатель: {PAYMENT_DETAILS['recipient_name']}\n"
                f"ИНН: {PAYMENT_DETAILS['inn']}\n"
                f"БИК: {PAYMENT_DETAILS['bic']}\n"
                f"Счет: {PAYMENT_DETAILS['account']}\n"
                f"Назначение: {payment_purpose}\n\n"
                f"Отсканируй QR-код для оплаты или используй реквизиты выше.\n\n"
                f"После оплаты отправьте фото чека в этот чат."
            )
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.answer()
                try:
                    await context.bot.send_photo(
                        chat_id=tg_user_id,
                        photo=qr_buf,
                        caption=payment_text,
                        parse_mode='Markdown'
                    )
                except Exception as e:
                    try:
                        payment_text_plain = payment_text.replace('*', '')
                        await context.bot.send_photo(
                            chat_id=tg_user_id,
                            photo=qr_buf,
                            caption=payment_text_plain
                        )
                    except Exception as e2:
                        await update.callback_query.message.reply_text(
                            payment_text.replace('*', '') + "\n\nНе удалось отправить QR-код. Используй реквизиты выше."
                        )
            else:
                await update.message.reply_text(payment_text, parse_mode='Markdown')
                try:
                    await context.bot.send_photo(
                        chat_id=tg_user_id,
                        photo=qr_buf,
                        caption="QR-код для оплаты"
                    )
                except Exception as e:
                    pass
        except Exception as qr_error:
            import traceback
            traceback.print_exc()
            room_parts = []
            if room_building:
                room_parts.append(str(room_building))
            if room_entrance:
                room_parts.append(str(room_entrance))
            if room_number_display:
                room_parts.append(str(room_number_display))
            
            if room_parts:
                room_info = "-".join(room_parts)
            else:
                room_info = "не указана"
            
            payment_text = (
                f"*Оплата проживания*\n\n"
                f"Студент: {student_fio}\n"
                f"Комната: {room_info}\n"
                f"Количество проживающих: {room_occupancy}\n"
                f"*Сумма к оплате: {payment_amount} руб.*\n\n"
                f"*Реквизиты для оплаты:*\n"
                f"Банк: {PAYMENT_DETAILS['bank_name']}\n"
                f"Получатель: {PAYMENT_DETAILS['recipient_name']}\n"
                f"ИНН: {PAYMENT_DETAILS['inn']}\n"
                f"БИК: {PAYMENT_DETAILS['bic']}\n"
                f"Счет: {PAYMENT_DETAILS['account']}\n"
                f"Назначение: {payment_purpose}\n\n"
                f"Не удалось сгенерировать QR-код. Используй реквизиты выше для оплаты.\n\n"
                f"После оплаты отправь фото чека в этот чат."
            )
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.answer("QR-код не сгенерирован, отправлены реквизиты", show_alert=True)
                await update.callback_query.message.reply_text(payment_text, parse_mode='Markdown')
            else:
                await update.message.reply_text(payment_text, parse_mode='Markdown')
    except Exception as e:
        import traceback
        traceback.print_exc()
        error_msg = f"Произошла ошибка: {str(e)}"
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer(f"{error_msg}", show_alert=True)
        else:
            await update.message.reply_text(f"{error_msg}")

async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text.strip()
    tg_user_id = update.effective_user.id
    tg_username = update.effective_user.username or ''
    
    registration_step = context.user_data.get('registration_step')
    
    if registration_step == 'building':
        if not text or not text.strip():
            await update.message.reply_text("Пожалуйста, введи номер корпуса")
            return
        
        context.user_data['building'] = text.strip()
        context.user_data['registration_step'] = 'entrance'
        await update.message.reply_text(
            f"Корпус: {text.strip()}\n\n"
            "Пожалуйста, отправьте номер подъезда:"
        )
    
    elif registration_step == 'entrance':
        if not text or not text.strip():
            await update.message.reply_text("Пожалуйста, введи номер подъезда")
            return
        
        context.user_data['entrance'] = text.strip()
        context.user_data['registration_step'] = 'room_number'
        await update.message.reply_text(
            f"Корпус: {context.user_data['building']}\n"
            f"Подъезд: {text.strip()}\n\n"
            "Пожалуйста, отправьте номер комнаты:"
        )
    
    elif registration_step == 'room_number':
        if not text or not text.strip():
            await update.message.reply_text("Пожалуйста, введи номер комнаты")
            return
        
        context.user_data['room_number'] = text.strip()
        context.user_data['registration_step'] = 'fio'
        await update.message.reply_text(
            f"Корпус: {context.user_data['building']}\n"
            f"Подъезд: {context.user_data['entrance']}\n"
            f"Номер: {text.strip()}\n\n"
            "Пожалуйста, отправьте ваше ФИО:"
        )
    
    elif registration_step == 'fio':
        if not text or not text.strip():
            await update.message.reply_text("Пожалуйста, введи свое ФИО")
            return
        
        context.user_data['fio'] = text.strip()
        context.user_data['registration_step'] = 'phone'
        await update.message.reply_text(
            f"Корпус: {context.user_data['building']}\n"
            f"Подъезд: {context.user_data['entrance']}\n"
            f"Номер: {context.user_data['room_number']}\n"
            f"ФИО: {text.strip()}\n\n"
            "Пожалуйста, введите ваш номер телефона:"
        )
    
    elif registration_step == 'phone':
        if not text or not text.strip():
            await update.message.reply_text("Пожалуйста, введи свой номер телефона")
            return
        
        student = find_student_by_room_and_info(
            context.user_data['building'],
            context.user_data['entrance'],
            context.user_data['room_number'],
            context.user_data['fio'],
            text.strip()
        )
        
        if student:
            if not check_username_match(student['id'], tg_username):
                await update.message.reply_text(
                    "Это не твой Telegram аккаунт.\n\n"
                    "Если ты сменил username, обратись к воспитателям для обновления данных."
                )
                context.user_data.pop('registration_step', None)
                context.user_data.pop('building', None)
                context.user_data.pop('entrance', None)
                context.user_data.pop('room_number', None)
                context.user_data.pop('fio', None)
                return
            
            register_tg_user(tg_user_id, tg_username, student['id'])
            context.user_data.pop('registration_step', None)
            context.user_data.pop('building', None)
            context.user_data.pop('entrance', None)
            context.user_data.pop('room_number', None)
            context.user_data.pop('fio', None)
            
            await update.message.reply_text(
                f"Отлично! Ты успешно зарегистрирован!\n\n"
                f"ФИО: {student['fio']}\n"
                f"Телефон: {student['phone']}\n"
                f"Группа: {student['group_name']}\n"
                f"Комната: {student['building']}-{student['entrance']}-{student['room_number']}\n\n"
            )
            await show_main_menu(update, context)
        else:
            await update.message.reply_text(
                "Не удалось найти тебя в базе данных по указанным данным.\n\n"
                "Проверь правильность введенных данных:\n"
                f"• Корпус: {context.user_data['building']}\n"
                f"• Подъезд: {context.user_data['entrance']}\n"
                f"• Номер: {context.user_data['room_number']}\n"
                f"• ФИО: {context.user_data['fio']}\n"
                f"• Телефон: {text.strip()}\n\n"
                "Начни регистрацию заново командой /start"
            )
            context.user_data.pop('registration_step', None)
            context.user_data.pop('building', None)
            context.user_data.pop('entrance', None)
            context.user_data.pop('room_number', None)
            context.user_data.pop('fio', None)
    
    elif registration_step == 'local_fio':
        if not text or not text.strip():
            await update.message.reply_text("Пожалуйста, введи свое ФИО")
            return
        
        context.user_data['fio'] = text.strip()
        context.user_data['registration_step'] = 'local_phone'
        await update.message.reply_text(
            f"ФИО: {text.strip()}\n\n"
            "Пожалуйста, введите ваш номер телефона:"
        )
    
    elif registration_step == 'local_phone':
        if not text or not text.strip():
            await update.message.reply_text("Пожалуйста, введи свой номер телефона")
            return
        
        student = find_local_student_by_info(
            context.user_data['fio'],
            text.strip()
        )
        
        if student:
            if not check_username_match(student['id'], tg_username):
                await update.message.reply_text(
                    "Это не твой Telegram аккаунт.\n\n"
                    "Если ты сменил username, обратись к воспитателям для обновления данных."
                )
                context.user_data.pop('registration_step', None)
                context.user_data.pop('is_local', None)
                context.user_data.pop('fio', None)
                return
            
            register_tg_user(tg_user_id, tg_username, student['id'])
            context.user_data.pop('registration_step', None)
            context.user_data.pop('is_local', None)
            context.user_data.pop('fio', None)
            
            await update.message.reply_text(
                f"Отлично! Ты успешно зарегистрирован как местный студент!\n\n"
                f"ФИО: {student['fio']}\n"
                f"Телефон: {student['phone']}\n"
                f"Группа: {student['group_name']}\n\n"
            )
            await show_main_menu(update, context)
        else:
            # Отладочная информация - проверяем, есть ли вообще местные студенты в базе
            with get_db() as conn:
                cursor = conn.cursor()
                cursor.execute("PRAGMA table_info(employees)")
                columns = [col[1] for col in cursor.fetchall()]
                has_is_local = 'is_local' in columns
                
                # Проверяем количество местных студентов
                if has_is_local:
                    cursor.execute("SELECT COUNT(*) as cnt FROM employees WHERE is_local = 1")
                else:
                    cursor.execute("""
                        SELECT COUNT(*) as cnt FROM employees 
                        WHERE (COALESCE(building, '') = '' OR building = '0')
                        AND (COALESCE(entrance, '') = '' OR entrance = '0')
                        AND (COALESCE(room_number, '') = '' OR room_number = '0')
                    """)
                local_count = cursor.fetchone()['cnt']
            
            error_msg = (
                "Не удалось найти тебя в базе данных по указанным данным.\n\n"
                "Проверь правильность введенных данных:\n"
                f"• ФИО: {context.user_data['fio']}\n"
                f"• Телефон: {text.strip()}\n\n"
            )
            
            if local_count == 0:
                error_msg += "⚠️ В базе данных не найдено местных студентов.\n"
                error_msg += "Обратись к администратору для добавления твоих данных.\n\n"
            else:
                error_msg += "💡 Убедись, что:\n"
                error_msg += "• ФИО введено полностью и правильно\n"
                error_msg += "• Номер телефона соответствует тому, что указан в базе\n\n"
            
            error_msg += "Начни регистрацию заново командой /start"
            
            await update.message.reply_text(error_msg)
            context.user_data.pop('registration_step', None)
            context.user_data.pop('is_local', None)
            context.user_data.pop('fio', None)
    
    elif registration_step == 'maintenance_message':
        if not text or not text.strip():
            await update.message.reply_text("Пожалуйста, опиши проблему.")
            return
        
        tg_user = get_tg_user(tg_user_id)
        if not tg_user or not tg_user['employee_id']:
            await update.message.reply_text("Ты еще не зарегистрирован! Используй /start для регистрации.")
            context.user_data.pop('registration_step', None)
            return
        
        employee_id = tg_user['employee_id']
        with get_db() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT fio, building, entrance, room_number FROM employees WHERE id = ?", (employee_id,))
            student = cursor.fetchone()
            
            if not student:
                await update.message.reply_text("Студент не найден в базе данных.")
                context.user_data.pop('registration_step', None)
                return
            
            import secrets
            request_number = f"MR-{secrets.token_hex(4).upper()}"
            created_at = datetime.now(MOSCOW_TZ).strftime('%Y-%m-%d %H:%M:%S')
            
            cursor.execute("""
                INSERT INTO maintenance_requests 
                (employee_id, tg_user_id, student_fio, building, entrance, room_number, message, status, request_number, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, 'new', ?, ?)
            """, (
                employee_id, tg_user_id, student['fio'] or '',
                student['building'] or '', student['entrance'] or '', student['room_number'] or '',
                text.strip(), request_number, created_at
            ))
            conn.commit()
        
        context.user_data.pop('registration_step', None)
        await update.message.reply_text(
            f"Отлично! Заявка успешно создана!\n\n"
            f"Описание: {text.strip()}\n\n"
            f"Твоя заявка будет рассмотрена администратором."
        )
        await show_main_menu(update, context)
    
    elif context.user_data.get('scan_event_step') == 'waiting_qr':
        event_id = context.user_data.get('selected_event_id')
        if not event_id:
            await update.message.reply_text("Мероприятие не выбрано. Начни заново.")
            context.user_data.pop('scan_event_step', None)
            return
        
        qr_code = text.strip()
        if not qr_code:
            await update.message.reply_text("Пожалуйста, отсканируй qr.")
            return
        
        await process_event_qr_code(update, context, tg_user_id, event_id, qr_code)
    
    else:
        tg_user = get_tg_user(tg_user_id)
        if tg_user and tg_user['employee_id']:
            # Если пользователь не в процессе регистрации или создания заявки, показываем меню
            await show_main_menu(update, context)
        else:
            context.user_data['registration_step'] = 'building'
            await update.message.reply_text(
                "Привет! Я твой помощник в студенческом кампусе.\n\n"
                "Для регистрации мне нужна следующая информация:\n\n"
                "Пожалуйста, отправь номер корпуса:"
            )

async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tg_user_id = update.effective_user.id
    tg_user = get_tg_user(tg_user_id)
    
    if not tg_user or not tg_user['employee_id']:
        await update.message.reply_text("Ты еще не зарегистрирован! Используй /start для регистрации.")
        return
    
    if context.user_data.get('scan_event_step') == 'waiting_qr' and context.user_data.get('selected_event_id'):
        await process_event_qr_photo(update, context, tg_user_id)
        return
    
    photo = update.message.photo
    if not photo:
        await update.message.reply_text("Не удалось получить фото. Попробуй еще раз.")
        return
    
    file = await context.bot.get_file(photo[-1].file_id)
    
    try:
        receipt_folder = 'static/receipts'
        os.makedirs(receipt_folder, exist_ok=True)
        
        filename = f"{tg_user['employee_id']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.jpg"
        filepath = os.path.join(receipt_folder, filename)
        
        await file.download_to_drive(filepath)
        
        payment_date = datetime.now(MOSCOW_TZ).strftime('%Y-%m-%d %H:%M:%S')
        now = datetime.now(MOSCOW_TZ)
        if now.month == 1:
            payment_month = f"{now.year - 1}-12"
        else:
            payment_month = f"{now.year}-{now.month - 1:02d}"
        
        with get_db() as conn:
            conn.execute('''
                INSERT INTO dormitory_payments (employee_id, tg_user_id, receipt_file, status, payment_date, payment_month, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            ''', (tg_user['employee_id'], tg_user_id, filepath, 'pending', payment_date, payment_month, payment_date))
            conn.commit()
        
            await update.message.reply_text(
            "Отлично! Чек успешно загружен!\n\n"
            "Ожидай проверки аудитом. Ты получишь уведомление о результатах проверки."
            )
    except Exception as e:
        import traceback
        traceback.print_exc()
        await update.message.reply_text("Произошла ошибка при обработке фото. Попробуй позже.")

async def process_event_qr_photo(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    event_id = context.user_data.get('selected_event_id')
    if not event_id:
        await update.message.reply_text("Мероприятие не выбрано. Начни заново.")
        return
    
    photo = update.message.photo
    if not photo:
        await update.message.reply_text("Упс! Не удалось получить фото. Попробуй еще раз.")
        return
    
    try:
        file = await context.bot.get_file(photo[-1].file_id)
        photo_buffer = io.BytesIO()
        await file.download_to_memory(photo_buffer)
        photo_buffer.seek(0)
        
        qr_code = None
        try:
            from PIL import Image
            from pyzbar import pyzbar
            
            img = Image.open(photo_buffer)
            decoded_objects = pyzbar.decode(img)
            
            if decoded_objects:
                qr_code = decoded_objects[0].data.decode('utf-8')
        except ImportError:
            try:
                import cv2
                import numpy as np
                
                photo_buffer.seek(0)
                img_bytes = photo_buffer.read()
                img_array = np.frombuffer(img_bytes, dtype=np.uint8)
                img = cv2.imdecode(img_array, cv2.IMREAD_COLOR)
                
                if img is not None:
                    detector = cv2.QRCodeDetector()
                    retval, decoded_info, points, straight_qrcode = detector.detectAndDecodeMulti(img)
                    
                    if retval and decoded_info and len(decoded_info) > 0:
                        qr_code = decoded_info[0]
            except ImportError:
                await update.message.reply_text(
                    "Для распознавания QR-кодов из фото нужно установить библиотеку.\n\n"
                    "Установи одну из:\n"
                    "• pip install pyzbar pillow\n"
                    "• pip install opencv-python pillow\n\n"
                )
                return
            except Exception as e:
                pass
        except Exception as e:
            await update.message.reply_text(
                f"Упс! Не удалось распознать QR-код.\n"
            )
            return
        
        if 't.me' in qr_code or 'tg://' in qr_code:
            token = None
            if 'start=event_' in qr_code:
                token = qr_code.split('start=event_')[1].split('&')[0].split('?')[0]
            elif '/event_' in qr_code:
                token = qr_code.split('/event_')[1].split('?')[0].split('&')[0]
            
            if token:
                try:
                    bot_info = await context.bot.get_me()
                    bot_username = bot_info.username
                    telegram_link = f"https://t.me/{bot_username}?start=event_{token}"
                except Exception:
                    telegram_link = qr_code
                
                message = (
                    "Отлично! QR-код распознан!\n\n"
                    "Нажми на ссылку ниже для регистрации студента:\n\n"
                    f"{telegram_link}"
                )
                
                keyboard = [
                    [InlineKeyboardButton("Открыть ссылку", url=telegram_link)],
                    [InlineKeyboardButton("Отмена", callback_data="back_to_menu")]
                ]
                reply_markup = InlineKeyboardMarkup(keyboard)
                
                await update.message.reply_text(message, reply_markup=reply_markup)
                return
        
        await process_event_qr_code(update, context, tg_user_id, event_id, qr_code)
        
    except Exception as e:
        import traceback
        traceback.print_exc()
        await update.message.reply_text("Упс! Что-то пошло не так при обработке фото. Попробуй позже.")

async def register_student_to_event(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id, event_id):
    employee_id = context.user_data.get('event_student_id')
    student_fio = context.user_data.get('event_student_fio')
    token = context.user_data.get('event_token')
    
    if not employee_id or not token:
        await update.callback_query.answer("Упс! Данные студента не найдены.", show_alert=True)
        return
    
    tg_username = update.effective_user.username or ''
    
    with get_db() as conn:
        cursor = conn.cursor()
        
        event = cursor.execute('SELECT id, organizers, location, event_date, event_time FROM events WHERE id = ?', (event_id,)).fetchone()
        if not event:
            await update.callback_query.answer("Мероприятие не найдено", show_alert=True)
            return
        
        is_global_org = cursor.execute(
            'SELECT id FROM event_organizers WHERE tg_username = ? OR tg_user_id = ?',
            (tg_username.lstrip('@'), tg_user_id)
        ).fetchone()
        
        if not is_global_org:
            organizers = event['organizers'] or ''
            if organizers:
                organizer_list = [o.strip() for o in organizers.split(',') if o.strip()]
                if tg_username.lstrip('@') not in organizer_list:
                    await update.callback_query.answer("У тебя нет прав организатора этого мероприятия", show_alert=True)
                    return
            else:
                await update.callback_query.answer("У тебя нет прав организатора этого мероприятия", show_alert=True)
                return
        
        existing = cursor.execute(
            'SELECT id FROM event_attendance WHERE event_id = ? AND employee_id = ?',
            (event_id, employee_id)
        ).fetchone()
        
        if existing:
            await update.callback_query.answer(f"Студент {student_fio} уже отмечен на этом мероприятии", show_alert=True)
            return
        scanned_at = datetime.now(MOSCOW_TZ).strftime('%Y-%m-%d %H:%M:%S')
        cursor.execute('''
            INSERT INTO event_attendance (event_id, employee_id, student_fio, scanned_by_tg_id, scanned_by_username, scanned_at)
            VALUES (?, ?, ?, ?, ?, ?)
        ''', (event_id, employee_id, student_fio, tg_user_id, tg_username.lstrip('@'), scanned_at))
        conn.commit()
        
        event_info = f"{event['location']} ({event['event_date']} {event['event_time']})"
        await update.callback_query.answer(f"Отлично! {student_fio} зарегистрирован!", show_alert=False)
        
        message = (
            f"Готово! Студент успешно зарегистрирован!\n\n"
            f"{student_fio}\n"
            f"{event_info}"
        )
        
        keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
        context.user_data.pop('event_token', None)
        context.user_data.pop('event_student_id', None)
        context.user_data.pop('event_student_fio', None)
        context.user_data.pop('event_step', None)

async def process_event_qr_code(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id, event_id, qr_code):
    tg_username = update.effective_user.username or ''
    
    token = None
    employee_id = None
    
    if 't.me' in qr_code or 'tg://' in qr_code:
        if 'start=event_' in qr_code:
            token = qr_code.split('start=event_')[1].split('&')[0].split('?')[0]
        elif '/event_' in qr_code:
            token = qr_code.split('/event_')[1].split('?')[0].split('&')[0]
    elif qr_code.startswith('STUDENT_QR_'):
        try:
            employee_id = int(qr_code.replace('STUDENT_QR_', ''))
        except ValueError:
            await update.message.reply_text("Упс! Неверный формат ID студента в QR-коде.")
            return
    elif len(qr_code) > 20:
        token = qr_code
    if token:
        with get_db() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT employee_id, expires_at FROM student_event_tokens 
                WHERE token = ?
            """, (token,))
            token_data = cursor.fetchone()
            
            if not token_data:
                await update.message.reply_text(
                    "Упс! Неверный токен QR-кода.\n"
                    "Убедись, что сканируешь актуальный QR-код студента."
                )
                return
            
            expires_at = datetime.fromisoformat(token_data['expires_at'])
            if expires_at.tzinfo is None:
                expires_at = expires_at.replace(tzinfo=MOSCOW_TZ)
            else:
                expires_at = expires_at.astimezone(MOSCOW_TZ)
            
            if datetime.now(MOSCOW_TZ) > expires_at:
                await update.message.reply_text(
                    "Ой! Токен QR-кода истек.\n"
                    "Студенту нужно получить новый QR-код."
                )
                return
            
            employee_id = token_data['employee_id']
    
    if not employee_id:
        await update.message.reply_text("Упс! Не удалось определить студента из QR-кода.")
        return
    await process_event_qr_direct(update, context, tg_user_id, event_id, employee_id, tg_username)

async def process_event_qr_direct(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id, event_id, employee_id, tg_username):
    with get_db() as conn:
        cursor = conn.cursor()
        
        event = cursor.execute('SELECT id, organizers FROM events WHERE id = ?', (event_id,)).fetchone()
        if not event:
            await update.message.reply_text("Мероприятие не найдено.")
            return
        
        is_global_org = cursor.execute(
            'SELECT id FROM event_organizers WHERE tg_username = ? OR tg_user_id = ?',
            (tg_username.lstrip('@'), tg_user_id)
        ).fetchone()
        
        if not is_global_org:
            organizers = event['organizers'] or ''
            if organizers:
                organizer_list = [o.strip() for o in organizers.split(',') if o.strip()]
                if tg_username.lstrip('@') not in organizer_list:
                    await update.message.reply_text("У тебя нет прав организатора этого мероприятия.")
                    return
            else:
                await update.message.reply_text("У тебя нет прав организатора этого мероприятия.")
                return
        
        student = cursor.execute('SELECT id, fio FROM employees WHERE id = ?', (employee_id,)).fetchone()
        if not student:
            await update.message.reply_text("Студент не найден в базе данных.")
            return
        
        existing = cursor.execute(
            'SELECT id FROM event_attendance WHERE event_id = ? AND employee_id = ?',
            (event_id, employee_id)
        ).fetchone()
        
        if existing:
            await update.message.reply_text(
                f"Студент {student['fio']} уже отмечен на этом мероприятии."
            )
            return
        
        scanned_at = datetime.now(MOSCOW_TZ).strftime('%Y-%m-%d %H:%M:%S')
        cursor.execute('''
            INSERT INTO event_attendance (event_id, employee_id, student_fio, scanned_by_tg_id, scanned_by_username, scanned_at)
            VALUES (?, ?, ?, ?, ?, ?)
        ''', (event_id, employee_id, student['fio'], tg_user_id, tg_username.lstrip('@'), scanned_at))
        conn.commit()
        
        await update.message.reply_text(
            f"Студент успешно отмечен!\n\n"
            f"{student['fio']}\n"
            f"Мероприятие: #{event_id}"
        )
        
        context.user_data.pop('selected_event_id', None)
        context.user_data.pop('scan_event_step', None)

def generate_student_event_token(employee_id):
    with get_db() as conn:
        cursor = conn.cursor()
        
        try:
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS student_event_tokens (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    employee_id INTEGER NOT NULL,
                    token TEXT UNIQUE NOT NULL,
                    created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                    expires_at TEXT NOT NULL,
                    FOREIGN KEY (employee_id) REFERENCES employees(id) ON DELETE CASCADE
                )
            """)
            conn.commit()
        except sqlite3.OperationalError:
            pass
        
        try:
            cursor.execute("""
                SELECT token FROM student_event_tokens 
                WHERE employee_id = ? AND expires_at > ?
            """, (employee_id, datetime.now(MOSCOW_TZ).isoformat()))
            existing = cursor.fetchone()
            
            if existing:
                return existing['token']
        except sqlite3.OperationalError:
            pass
        
        token = secrets.token_urlsafe(32)
        expires_at = datetime.now(MOSCOW_TZ) + timedelta(days=365)
        
        try:
            cursor.execute("""
                INSERT INTO student_event_tokens (employee_id, token, expires_at)
                VALUES (?, ?, ?)
            """, (employee_id, token, expires_at.isoformat()))
            conn.commit()
        except sqlite3.IntegrityError:
            token = secrets.token_urlsafe(32)
            cursor.execute("""
                INSERT INTO student_event_tokens (employee_id, token, expires_at)
                VALUES (?, ?, ?)
            """, (employee_id, token, expires_at.isoformat()))
            conn.commit()
        
        return token

async def generate_event_qr(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Ты еще не зарегистрирован! Используй /start для регистрации.", show_alert=True)
        return
    
    employee_id = tg_user['employee_id']
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT fio FROM employees WHERE id = ?", (employee_id,))
        student = cursor.fetchone()
    
    if not student:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Студент не найден", show_alert=True)
        return
    
    token = generate_student_event_token(employee_id)
    
    try:
        bot_info = await context.bot.get_me()
        bot_username = bot_info.username
    except Exception:
        bot_username = None
    
    if bot_username:
        telegram_link = f"https://t.me/{bot_username}?start=event_{token}"
    else:
        telegram_link = f"tg://resolve?domain={TELEGRAM_BOT_TOKEN.split(':')[0]}&start=event_{token}"
    
    qr_img = qrcode.make(telegram_link)
    qr_buf = io.BytesIO()
    qr_img.save(qr_buf, format='PNG')
    qr_buf.seek(0)
    
    caption = (
        f"Твой персональный QR-код для мероприятий\n\n"
        f"{student['fio']}\n\n"
        f"Покажи этот QR-код организатору мероприятия.\n"
        f"При сканировании откроется ссылка для регистрации."
    )
    
    keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    if hasattr(update, 'callback_query') and update.callback_query:
        await update.callback_query.answer()
        try:
            await context.bot.send_photo(
                chat_id=tg_user_id,
                photo=qr_buf,
                caption=caption,
                reply_markup=reply_markup
            )
        except Exception:
            await update.callback_query.message.reply_text(caption, reply_markup=reply_markup)
    elif hasattr(update, 'message') and update.message:
        try:
            await context.bot.send_photo(
                chat_id=tg_user_id,
                photo=qr_buf,
                caption=caption,
                reply_markup=reply_markup
            )
        except Exception:
            await update.message.reply_text(caption, reply_markup=reply_markup)

async def create_maintenance_request(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Ты еще не зарегистрирован! Используй /start для регистрации.", show_alert=True)
        return
    
    context.user_data['registration_step'] = 'maintenance_message'
    message = (
        "Подача заявки на обслуживание\n\n"
        "Опиши проблему или что нужно сделать:\n"
        "(Например: протекает кран, не работает розетка, нужна уборка и т.д.)"
    )
    
    keyboard = [[InlineKeyboardButton("Отмена", callback_data="back_to_menu")]]
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    if hasattr(update, 'callback_query') and update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
    elif hasattr(update, 'message') and update.message:
        await update.message.reply_text(message, reply_markup=reply_markup)

async def show_my_maintenance_requests(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Ты еще не зарегистрирован! Используй /start для регистрации.", show_alert=True)
        return
    
    employee_id = tg_user['employee_id']
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT id, message, status, request_number, created_at, admin_comment
            FROM maintenance_requests
            WHERE employee_id = ?
            ORDER BY created_at DESC
            LIMIT 10
        """, (employee_id,))
        requests = cursor.fetchall()
    
    if not requests:
        message = "У тебя пока нет заявок на обслуживание."
    else:
        message = "Твои заявки на обслуживание:\n\n"
        for req in requests:
            status_text = {
                'new': 'Новая',
                'in_progress': 'В работе',
                'closed': 'Закрыта',
                'completed': 'Выполнена'
            }.get(req['status'], req['status'])
            
            req_num = req['request_number'] or f"#{req['id']}"
            created = req['created_at'][:10] if req['created_at'] else ''
            message += f"{req_num} - {status_text}\n"
            message += f"   {req['message'][:50]}{'...' if len(req['message']) > 50 else ''}\n"
            if req['admin_comment']:
                message += f"   Комментарий: {req['admin_comment'][:50]}\n"
            message += f"   {created}\n\n"
    
    keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    if hasattr(update, 'callback_query') and update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
    elif hasattr(update, 'message') and update.message:
        await update.message.reply_text(message, reply_markup=reply_markup)

async def scan_event_qr(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Ты еще не зарегистрирован! Используй /start для регистрации.", show_alert=True)
        return
    
    tg_username = update.effective_user.username or ''
    if not tg_username:
        message = "Для сканирования QR-кодов нужно указать username в Telegram."
        keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer()
            await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
        return
    
    with get_db() as conn:
        cursor = conn.cursor()
        is_global_org = cursor.execute(
            "SELECT id FROM event_organizers WHERE tg_username = ? OR tg_user_id = ?",
            (tg_username.lstrip('@'), tg_user_id)
        ).fetchone()
        
        today = datetime.now(MOSCOW_TZ).date().isoformat()
        if is_global_org:
            events = cursor.execute("""
                SELECT id, location, description, event_date, event_time, organizers
                FROM events
                WHERE event_date >= ?
                ORDER BY event_date, event_time
                LIMIT 20
            """, (today,)).fetchall()
        else:
            events = cursor.execute("""
                SELECT id, location, description, event_date, event_time, organizers
                FROM events
                WHERE event_date >= ? AND (organizers LIKE ? OR organizers = '')
                ORDER BY event_date, event_time
                LIMIT 20
            """, (today, f'%{tg_username.lstrip("@")}%')).fetchall()
    
    if not events:
        message = (
            "Сканирование QR-кода на мероприятии\n\n"
            "У тебя нет активных мероприятий для сканирования."
        )
        keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer()
            await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
        return
    
    keyboard = []
    for event in events:
        event_date = event['event_date']
        event_time = event['event_time']
        location = event['location'][:30] if event['location'] else 'Без места'
        button_text = f"{event_date} {event_time} - {location}"
        keyboard.append([InlineKeyboardButton(button_text, callback_data=f"select_event_{event['id']}")])
    keyboard.append([InlineKeyboardButton("Отмена", callback_data="back_to_menu")])
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    message = (
        "Сканирование QR-кода на мероприятии\n\n"
        "Выбери мероприятие для сканирования:"
    )
    
    if hasattr(update, 'callback_query') and update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
    elif hasattr(update, 'message') and update.message:
        await update.message.reply_text(message, reply_markup=reply_markup)

async def show_duty_schedule(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Ты еще не зарегистрирован! Используй /start для регистрации.", show_alert=True)
        return
    
    employee_id = tg_user['employee_id']
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT ea.building, ea.entrance
            FROM elder_assignments ea
            JOIN employees e ON ea.employee_id = e.id
            WHERE e.id = ? AND e.is_entrance_elder = 1
        """, (employee_id,))
        elder_info = cursor.fetchone()
    
    if not elder_info:
        message = "Ты не назначен старостой подъезда."
        keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer()
            await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
        elif hasattr(update, 'message') and update.message:
            await update.message.reply_text(message, reply_markup=reply_markup)
        return
    
    building = elder_info['building']
    entrance = elder_info['entrance']
    
    now = datetime.now(MOSCOW_TZ)
    current_month = now.month
    current_year = now.year
    
    if current_month == 12:
        next_month = 1
        next_year = current_year + 1
    else:
        next_month = current_month + 1
        next_year = current_year
    
    month_names = ['', 'январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 
                   'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь']
    
    message = (
        f"График дежурств\n\n"
        f"Корпус: {building}\n"
        f"Подъезд: {entrance}\n\n"
        f"Выбери месяц для генерации графика:"
    )
    
    keyboard = [
        [InlineKeyboardButton(
            f"{month_names[current_month].capitalize()} {current_year} (текущий)",
            callback_data=f"generate_duty_current"
        )],
        [InlineKeyboardButton(
            f"{month_names[next_month].capitalize()} {next_year} (следующий)",
            callback_data=f"generate_duty_next"
        )],
        [InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]
    ]
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    if hasattr(update, 'callback_query') and update.callback_query:
        await update.callback_query.answer()
        await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
    elif hasattr(update, 'message') and update.message:
        await update.message.reply_text(message, reply_markup=reply_markup)

async def generate_duty_schedule(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id, target_month=None, target_year=None):
    import pandas as pd
    import io
    import calendar
    import json
    
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Ты еще не зарегистрирован!", show_alert=True)
        return
    
    employee_id = tg_user['employee_id']
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT ea.building, ea.entrance
            FROM elder_assignments ea
            JOIN employees e ON ea.employee_id = e.id
            WHERE e.id = ? AND e.is_entrance_elder = 1
        """, (employee_id,))
        elder_info = cursor.fetchone()
        
        if not elder_info:
            message = "Ты не назначен старостой подъезда."
            keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.answer()
                await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
            return
        
        building = elder_info['building']
        entrance = elder_info['entrance']
        
        now = datetime.now(MOSCOW_TZ)
        if target_month is None:
            target_month = now.month
        if target_year is None:
            target_year = now.year
        cursor.execute("""
            SELECT DISTINCT 
                COALESCE(r.room_number, e.room_number) as room_number,
                COALESCE(r.capacity, 
                    (SELECT COUNT(*) FROM employees e2 
                     WHERE e2.building = e.building 
                     AND e2.entrance = e.entrance 
                     AND e2.room_number = e.room_number), 1) as occupancy
            FROM employees e
            LEFT JOIN rooms r ON r.building = e.building 
                AND r.entrance = e.entrance 
                AND r.room_number = e.room_number
            WHERE e.building = ? AND e.entrance = ?
                AND e.room_number IS NOT NULL 
                AND e.room_number != ''
                AND (r.excluded_from_duty IS NULL OR r.excluded_from_duty = 0)
            ORDER BY e.room_number
        """, (building, entrance))
        
        rooms_data = cursor.fetchall()
        
        if not rooms_data:
            message = f"Нет комнат в подъезде {building}-{entrance}."
            keyboard = [[InlineKeyboardButton("Назад", callback_data="duty_schedule")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.answer()
                await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
            return
        
        rooms_1 = []
        rooms_2 = []
        rooms_other = []
        
        for room in rooms_data:
            room_num = room['room_number']
            occupancy = room['occupancy'] or 1
            
            if occupancy == 1:
                rooms_1.append(room_num)
            elif occupancy == 2:
                rooms_2.append(room_num)
            else:
                rooms_other.append(room_num)
        
        cursor.execute("""
            SELECT last_room_number, last_month, last_year
            FROM duty_schedule_progress
            WHERE building = ? AND entrance = ?
        """, (building, entrance))
        progress = cursor.fetchone()
        
        all_rooms = rooms_1 + rooms_2 + rooms_other
        start_index = 0
        
        if progress and progress['last_room_number']:
            last_room = progress['last_room_number']
            if last_room in all_rooms:
                start_index = (all_rooms.index(last_room) + 1) % len(all_rooms)
        
        duty_pairs = []
        rooms_1_idx = 0
        rooms_2_idx = 0
        rooms_other_idx = 0
        
        while rooms_1_idx < len(rooms_1) and rooms_2_idx < len(rooms_2):
            pair = f"{rooms_1[rooms_1_idx]}, {rooms_2[rooms_2_idx]}"
            duty_pairs.append(pair)
            rooms_1_idx += 1
            rooms_2_idx += 1
        
        while rooms_1_idx < len(rooms_1):
            duty_pairs.append(rooms_1[rooms_1_idx])
            rooms_1_idx += 1
        
        while rooms_2_idx < len(rooms_2):
            duty_pairs.append(rooms_2[rooms_2_idx])
            rooms_2_idx += 1
        
        for room in rooms_other:
            duty_pairs.append(room)
        
        if not duty_pairs:
            duty_pairs = all_rooms
        
        duty_pairs = duty_pairs[start_index:] + duty_pairs[:start_index]
        
        days_in_month = calendar.monthrange(target_year, target_month)[1]
        
        first_day = datetime(target_year, target_month, 1)
        first_weekday = first_day.weekday()
        
        schedule_data = {}
        current_pair_idx = 0
        
        for day in range(1, days_in_month + 1):
            date_str = f"{day:02d}.{target_month:02d}.{target_year}"
            if len(duty_pairs) > 0:
                room_assignment = duty_pairs[current_pair_idx % len(duty_pairs)]
                schedule_data[date_str] = room_assignment
                current_pair_idx += 1
        
        schedule_json = json.dumps(schedule_data, ensure_ascii=False)
        cursor.execute("""
            INSERT OR REPLACE INTO duty_schedules 
            (building, entrance, schedule_month, schedule_year, schedule_data, created_by)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (building, entrance, target_month, target_year, schedule_json, employee_id))
        
        last_assignment = schedule_data.get(f"{days_in_month:02d}.{target_month:02d}.{target_year}", "")
        if last_assignment:
            if ',' in last_assignment:
                last_room = last_assignment.split(',')[-1].strip()
            else:
                last_room = last_assignment.strip()
            
            cursor.execute("""
                INSERT OR REPLACE INTO duty_schedule_progress 
                (building, entrance, last_room_index, last_room_number, last_month, last_year, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, datetime('now'))
            """, (building, entrance, current_pair_idx, last_room, target_month, target_year))
        
        conn.commit()
        
        month_names = ['', 'январь', 'февраль', 'март', 'апрель', 'май', 'июнь', 
                       'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь']
        
        excel_data = []
        weekdays = ['ПН', 'ВТ', 'СР', 'ЧТ', 'ПТ', 'СБ', 'ВС']
        
        header_row = {'День недели': 'День недели'}
        for day in weekdays:
            header_row[day] = day
        excel_data.append(header_row)
        
        instructions = [
            'Кухня: проверка чистоты (столы, мойка, плитка); выбросить мусор',
            'Постирочная: постиранное белье выгружено, чистота в помещении',
            'Гладильная: глад. доски на месте на каждом этаже, утюги выключены, покрытие чистое',
            'Коворкинги и комнаты самоподготовки: чистота, порядок в комнате',
            'Лифт: чистота'
        ]
        
        for instruction in instructions:
            instruction_row = {'День недели': 'Инструкция'}
            for day in weekdays:
                instruction_row[day] = instruction
            excel_data.append(instruction_row)
        
        empty_row = {'День недели': ''}
        for day in weekdays:
            empty_row[day] = ''
        excel_data.append(empty_row)
        
        day_num = 1
        current_pair_idx = 0
        
        for week in range(6):
            week_row = {'День недели': ''}
            week_has_days = False
            
            for weekday_idx in range(7):
                if day_num > days_in_month:
                    week_row[weekdays[weekday_idx]] = ''
                elif week == 0 and weekday_idx < first_weekday:
                    week_row[weekdays[weekday_idx]] = ''
                else:
                    date_str = f"{day_num:02d}.{target_month:02d}.{target_year}"
                    
                    if len(duty_pairs) > 0:
                        room_assignment = duty_pairs[current_pair_idx % len(duty_pairs)]
                        week_row[weekdays[weekday_idx]] = f"{date_str}\n{room_assignment}"
                        current_pair_idx += 1
                    else:
                        week_row[weekdays[weekday_idx]] = date_str
                    
                    day_num += 1
                    week_has_days = True
            
            if week_has_days or week == 0:
                excel_data.append(week_row)
            
            if day_num > days_in_month:
                break
        
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, Border, Side
        from openpyxl.utils import get_column_letter

        buf = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = 'График дежурств'

        thin = Side(style='thin', color='000000')
        medium = Side(style='medium', color='000000')
        header_border = Border(top=medium, left=medium, right=medium, bottom=medium)
        cell_border = Border(top=thin, left=thin, right=thin, bottom=thin)

        ws.column_dimensions['A'].width = 30
        for col in range(2, 9):
            ws.column_dimensions[get_column_letter(col)].width = 15

        current_row = 1

        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=8)
        title_cell = ws.cell(row=current_row, column=1, value=f"График дежурств {building}-{entrance}")
        title_cell.font = Font(size=20, bold=True)
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        current_row += 2

        instruction_header_row = current_row
        ws.cell(row=instruction_header_row, column=1, value="Инструкция").font = Font(bold=True)
        for idx, instruction in enumerate(instructions, start=1):
            row_idx = instruction_header_row + idx
            ws.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=8)
            cell = ws.cell(row=row_idx, column=2, value=instruction)
            cell.alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
        current_row = instruction_header_row + len(instructions) + 2

        weekdays_full = ['ПН', 'ВТ', 'СР', 'ЧТ', 'ПТ', 'СБ', 'ВС']
        header_row = current_row
        ws.cell(row=header_row, column=1, value="")
        for idx, day in enumerate(weekdays_full, start=2):
            cell = ws.cell(row=header_row, column=idx, value=day)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = header_border
        current_row += 1

        day_num = 1
        first_day = datetime(target_year, target_month, 1)
        first_weekday = first_day.weekday()  # 0 = Monday
        days_in_month = calendar.monthrange(target_year, target_month)[1]

        while day_num <= days_in_month:
            date_row = current_row
            room_row = current_row + 1

            for weekday_idx in range(7):
                col = 2 + weekday_idx
                date_cell = ws.cell(row=date_row, column=col)
                room_cell = ws.cell(row=room_row, column=col)

                if current_row == header_row + 1 and weekday_idx < first_weekday:
                    date_cell.value = ""
                    room_cell.value = ""
                elif day_num <= days_in_month:
                    date_str = f"{day_num:02d}.{target_month:02d}.{target_year}"
                    date_cell.value = date_str
                    date_cell.alignment = Alignment(horizontal='center', vertical='center')
                    room_assignment = schedule_data.get(date_str, "")
                    room_cell.value = room_assignment
                    room_cell.font = Font(bold=True, size=14)
                    room_cell.alignment = Alignment(horizontal='center', vertical='center')
                    day_num += 1
                else:
                    date_cell.value = ""
                    room_cell.value = ""

                date_cell.border = cell_border
                room_cell.border = cell_border

            current_row += 2

        wb.save(buf)
        buf.seek(0)
        
        filename = f'График_дежурств_{building}_{entrance}_{month_names[target_month]}_{target_year}.xlsx'
        
        try:
            from telegram import InputFile
            document = InputFile(buf, filename=filename)
            
            caption = (
                f"График дежурств\n\n"
                f"Корпус: {building}\n"
                f"Подъезд: {entrance}\n"
                f"Месяц: {month_names[target_month].capitalize()} {target_year}\n\n"
                f"График успешно сгенерирован!"
            )
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.answer("Генерация файла...")
                await update.callback_query.message.reply_document(
                    document=document,
                    caption=caption
                )
            elif hasattr(update, 'message') and update.message:
                await update.message.reply_document(
                    document=document,
                    caption=caption
                )
            
            keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.message.reply_text("Файл отправлен!", reply_markup=reply_markup)
            elif hasattr(update, 'message') and update.message:
                await update.message.reply_text("Файл отправлен!", reply_markup=reply_markup)
                
        except Exception as e:
            error_message = f"Ошибка при отправке файла: {str(e)}"
            keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.message.reply_text(error_message, reply_markup=reply_markup)
            elif hasattr(update, 'message') and update.message:
                await update.message.reply_text(error_message, reply_markup=reply_markup)

async def assign_round(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer("Вы не зарегистрированы. Используйте /start для регистрации.", show_alert=True)
        return
    
    employee_id = tg_user['employee_id']
    with get_db() as conn:
        chief = conn.execute('''
            SELECT id, is_round_chief, fio
            FROM employees
            WHERE id = ?
        ''', (employee_id,)).fetchone()
        
        if not chief or not chief['is_round_chief']:
            message = "У тебя нет прав для назначения на обход.\n\nОбратись к администратору."
            keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.answer()
                await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
            elif hasattr(update, 'message') and update.message:
                await update.message.reply_text(message, reply_markup=reply_markup)
            return
        
        students = conn.execute('''
            SELECT id, fio, building, entrance, room_number, group_name
            FROM employees
            WHERE can_do_rounds = 1
            ORDER BY fio
            LIMIT 50
        ''').fetchall()
        
        if not students:
            message = (
                "Назначение на обход\n\n"
                "Нет студентов с правом обхода.\n"
                "Обратись к администратору для предоставления права обхода студентам."
            )
            keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.answer()
                await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
            elif hasattr(update, 'message') and update.message:
                await update.message.reply_text(message, reply_markup=reply_markup)
            return
        
        buildings_entrances = conn.execute('''
            SELECT DISTINCT building, entrance
            FROM employees
            WHERE building IS NOT NULL AND building != ''
            AND entrance IS NOT NULL AND entrance != ''
            ORDER BY building, entrance
        ''').fetchall()
        
        if not buildings_entrances:
            message = (
                "Назначение на обход\n\n"
                "Нет доступных корпусов и подъездов.\n"
                "Обратись к администратору."
            )
            keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
            reply_markup = InlineKeyboardMarkup(keyboard)
            
            if hasattr(update, 'callback_query') and update.callback_query:
                await update.callback_query.answer()
                await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
            elif hasattr(update, 'message') and update.message:
                await update.message.reply_text(message, reply_markup=reply_markup)
            return
        
        message = "Назначение на обход\n\nВыбери студента:\n\n"
        
        keyboard = []
        for student in students:
            room_info = ""
            if student['room_number']:
                room_info = f" (комн. {student['room_number']})"
            button_text = f"{student['fio']}{room_info}"
            if len(button_text) > 60:
                button_text = button_text[:57] + "..."
            keyboard.append([InlineKeyboardButton(
                button_text,
                callback_data=f"assign_round_student_{student['id']}"
            )])
        
        keyboard.append([InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")])
        reply_markup = InlineKeyboardMarkup(keyboard)
        
        context.user_data['assign_round_students'] = {s['id']: dict(s) for s in students}
        context.user_data['assign_round_buildings_entrances'] = [
            {'building': row['building'], 'entrance': row['entrance']}
            for row in buildings_entrances
        ]
        
        if hasattr(update, 'callback_query') and update.callback_query:
            await update.callback_query.answer()
            await update.callback_query.message.reply_text(message, reply_markup=reply_markup)
        elif hasattr(update, 'message') and update.message:
            await update.message.reply_text(message, reply_markup=reply_markup)

async def handle_assign_round_student(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id, student_id):
    students = context.user_data.get('assign_round_students', {})
    buildings_entrances = context.user_data.get('assign_round_buildings_entrances', [])
    
    if student_id not in students:
        await update.callback_query.answer("Студент не найден", show_alert=True)
        return
    
    student = students[student_id]
    
    if not buildings_entrances:
        await update.callback_query.answer("Нет доступных корпусов и подъездов", show_alert=True)
        return
    
    buildings_dict = {}
    for be in buildings_entrances:
        building = be['building']
        entrance = be['entrance']
        if building not in buildings_dict:
            buildings_dict[building] = []
        buildings_dict[building].append(entrance)
    
    message = f"Назначение на обход\n\n"
    message += f"Студент: {student['fio']}\n"
    if student['room_number']:
        message += f"Комната: {student['room_number']}\n"
    message += f"\nВыбери корпус и подъезд:\n\n"
    
    keyboard = []
    for building in sorted(buildings_dict.keys()):
        entrances = sorted(buildings_dict[building])
        for entrance in entrances:
            button_text = f"Корпус {building}, Подъезд {entrance}"
            keyboard.append([InlineKeyboardButton(
                button_text,
                callback_data=f"assign_round_building_{building}|{entrance}|{student_id}"
            )])
    
    keyboard.append([InlineKeyboardButton("Назад", callback_data="assign_round")])
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    await update.callback_query.answer()
    await update.callback_query.message.reply_text(message, reply_markup=reply_markup)

async def handle_assign_round_final(update: Update, context: ContextTypes.DEFAULT_TYPE, tg_user_id, student_id, building, entrance):
    tg_user = get_tg_user(tg_user_id)
    if not tg_user or not tg_user['employee_id']:
        await update.callback_query.answer("Ошибка: пользователь не найден", show_alert=True)
        return
    
    employee_id = tg_user['employee_id']
    
    with get_db() as conn:
        chief = conn.execute('''
            SELECT id, is_round_chief, fio
            FROM employees
            WHERE id = ?
        ''', (employee_id,)).fetchone()
        
        if not chief or not chief['is_round_chief']:
            await update.callback_query.answer("У тебя нет прав для назначения на обход", show_alert=True)
            return
        
        student = conn.execute('''
            SELECT id, fio, can_do_rounds 
            FROM employees 
            WHERE id = ?
        ''', (student_id,)).fetchone()
        
        if not student:
            await update.callback_query.answer("Студент не найден", show_alert=True)
            return
        
        if not student['can_do_rounds']:
            await update.callback_query.answer("У этого студента нет права обхода", show_alert=True)
            return
        
        try:
            conn.execute('''
                INSERT OR REPLACE INTO round_assignments 
                (student_id, building, entrance, assigned_by)
                VALUES (?, ?, ?, ?)
            ''', (student_id, building, entrance, chief['id']))
            conn.commit()
            
            message = (
                f"Назначение выполнено!\n\n"
                f"Студент: {student['fio']}\n"
                f"Корпус: {building}\n"
                f"Подъезд: {entrance}\n\n"
                f"Студент успешно назначен на обход."
            )
        except Exception as e:
            message = f"Ошибка при назначении: {str(e)}"
    
    keyboard = [[InlineKeyboardButton("Назад в меню", callback_data="back_to_menu")]]
    reply_markup = InlineKeyboardMarkup(keyboard)
    
    await update.callback_query.answer()
    await update.callback_query.message.reply_text(message, reply_markup=reply_markup)

def send_2fa_code_to_student(tg_user_id, employee_id):
    try:
        code = secrets.token_urlsafe(6)[:6].upper()
        expires_at = datetime.now(MOSCOW_TZ) + timedelta(minutes=10)
        
        with get_db() as conn:
            conn.execute("""
                INSERT INTO student_2fa_codes (employee_id, tg_user_id, code, created_at, expires_at)
                VALUES (?, ?, ?, ?, ?)
            """, (
                employee_id, tg_user_id, code,
                datetime.now(MOSCOW_TZ).isoformat(),
                expires_at.isoformat()
            ))
            conn.commit()
        
        try:
            import asyncio
            from telegram import Bot
            
            bot_token = TELEGRAM_BOT_TOKEN
            bot = Bot(token=bot_token)
            message = (
                f"Твой код для входа в личный кабинет:\n\n"
                f"||{code}||\n\n"
                f"Код действителен 10 минут"
            )
            
            try:
                loop = asyncio.get_event_loop()
                if loop.is_running():
                    asyncio.create_task(bot.send_message(chat_id=tg_user_id, text=message, parse_mode='Markdown'))
                else:
                    loop.run_until_complete(bot.send_message(chat_id=tg_user_id, text=message, parse_mode='Markdown'))
            except RuntimeError:
                asyncio.run(bot.send_message(chat_id=tg_user_id, text=message, parse_mode='Markdown'))
        except Exception as e:
            try:
                import asyncio
                from telegram import Bot
                bot_token = TELEGRAM_BOT_TOKEN
                bot = Bot(token=bot_token)
                message = (
                    f"Твой код для входа в личный кабинет:\n\n"
                    f"{code}\n\n"
                    f"Код действителен 10 минут"
                )
                try:
                    loop = asyncio.get_event_loop()
                    if loop.is_running():
                        asyncio.create_task(bot.send_message(chat_id=tg_user_id, text=message))
                    else:
                        loop.run_until_complete(bot.send_message(chat_id=tg_user_id, text=message))
                except RuntimeError:
                    asyncio.run(bot.send_message(chat_id=tg_user_id, text=message))
            except Exception:
                pass
        
        return code
    except Exception as e:
        import traceback
        traceback.print_exc()
        return None

async def event_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    tg_user_id = update.effective_user.id
    tg_username = update.effective_user.username or ''
    
    if not context.args or len(context.args) == 0:
        await update.message.reply_text(
            "Неверный формат команды.\n"
            "Используй: /event <токен>\n"
            "Или отсканируй QR-код студента."
        )
        return
    
    token = context.args[0]
    await handle_event_token(update, context, tg_user_id, token)

def main():
    import sys
    import os
    
    if not TELEGRAM_AVAILABLE:
        return
    
    if not TELEGRAM_BOT_TOKEN:
        print("TELEGRAM_BOT_TOKEN не установлен. Бот не будет запущен.")
        return
    
    try:
        import subprocess
        result = subprocess.run(
            ['ps', 'aux'],
            capture_output=True,
            text=True
        )
        if result.returncode == 0:
            lines = result.stdout.split('\n')
            telegram_bot_processes = [
                line for line in lines 
                if 'telegram_bot.py' in line and str(os.getpid()) not in line
            ]
            if telegram_bot_processes:
                print("ВНИМАНИЕ: Обнаружены другие запущенные экземпляры бота:")
                for proc in telegram_bot_processes:
                    print(f"   {proc}")
                print("\nЭто может вызвать ошибку 'Conflict: terminated by other getUpdates request'")
                print("Рекомендуется остановить другие процессы перед запуском.")
                print("   Для остановки используйте: kill <PID>")
                response = input("\nПродолжить запуск? (y/n): ")
                if response.lower() != 'y':
                    print("Запуск отменен.")
                    sys.exit(0)
    except Exception as e:
        pass
    
    print("Запуск Telegram бота...")
    
    async def post_init(application: Application) -> None:
        """Инициализация после создания приложения - настройка команд бота"""
        try:
            commands = [
                BotCommand("menu", "Главное меню"),
            ]
            await application.bot.set_my_commands(commands)
            print("Боковое меню настроено!")
        except Exception as e:
            print(f"Предупреждение: Не удалось установить команды бота: {e}")
            print("Бот продолжит работу, но команды могут быть недоступны в меню.")
    
    application = Application.builder().token(TELEGRAM_BOT_TOKEN).post_init(post_init).build()
    
    application.add_handler(CommandHandler("start", start))
    application.add_handler(CommandHandler("event", event_command))
    application.add_handler(CommandHandler("menu", menu_command))
    application.add_handler(CommandHandler("qr", qr_command))
    application.add_handler(CallbackQueryHandler(button_handler))
    application.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
    
    application.add_error_handler(error_handler)
    
    job_queue = application.job_queue
    
    job_queue.run_daily(
        send_bed_linen_reminders,
        time=dt_time(hour=7, minute=0),
        name="bed_linen_reminders"
    )
    
    print("Бот готов к работе!")
    
    try:
        application.run_polling(
            allowed_updates=Update.ALL_TYPES,
            drop_pending_updates=False
        )
    except Conflict as e:
        print("\nОШИБКА: Обнаружен конфликт с другим экземпляром бота!")
        print(f"   {str(e)}")
        print("\n   Решение:")
        print("   1. Найдите другие запущенные процессы: ps aux | grep telegram_bot")
        print("   2. Остановите их: kill <PID>")
        print("   3. Запустите бот снова")
        sys.exit(1)
    except NetworkError as e:
        print("\nОШИБКА: Проблема с подключением к Telegram API!")
        print(f"   {str(e)}")
        print("\n   Возможные причины:")
        print("   1. Проблемы с интернет-соединением")
        print("   2. Блокировка Telegram API файрволом/прокси")
        print("   3. Неверный токен бота")
        print("   4. Проблемы с SSL/TLS соединением")
        print("\n   Проверьте:")
        print("   - Интернет-соединение")
        print("   - Настройки прокси/файрвола")
        print("   - Правильность TELEGRAM_BOT_TOKEN в config.py")
        sys.exit(1)

if __name__ == '__main__':
    main()

